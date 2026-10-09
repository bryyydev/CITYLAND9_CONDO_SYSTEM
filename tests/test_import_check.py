"""Dry-run ("Check only") validator for data-migration workbooks (app/services/import_check.py,
CITYLAND9_IMPORT_SAFEGUARDS.md). Synthetic workbooks only. The check must never write."""
import hashlib
import io
from datetime import date, datetime

import pytest
from openpyxl import Workbook
from sqlalchemy import text

from conftest import PASSWORD, api, login

from app.services import import_check as ic

CONFIG = {
    "source_system": {"code": "synthetic-legacy", "name": "Synthetic legacy system"},
    "properties": {
        "unit_no_confirmed": True,
        "group_mapping": {
            "RESIDENTIAL": {"approved": True},
            "2F": {"approved": True, "floor": "2F"},
            "PARKING": {"approved": False},
        },
        "people": {"one_person_per_cell_confirmed": False},
    },
}
HEADERS = ["ID", "Name", "Size", "Group", "Owner", "Tenant"]


def xlsx(sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def check(app_module, sheets, config=CONFIG, roundtrip="auto"):
    with app_module.app.app_context():
        return ic.check_workbook(xlsx(sheets), "synthetic.xlsx", legacy=app_module, config=config, roundtrip=roundtrip)


def codes_at(report, row, column=None):
    return {i["code"] for i in report["issues"] if i["row"] == row and (column is None or i["column"] == column)}


def snapshot(app_module):
    """A digest of every row of every table: equal before/after = nothing was written."""
    digest = hashlib.sha256()
    with app_module.app.app_context():
        conn = app_module.db.session.connection()
        for table in sorted(app_module.db.metadata.tables):
            for row in conn.execute(text(f'SELECT * FROM "{table}" ORDER BY 1')).fetchall():
                digest.update(repr((table, tuple(row))).encode())
        app_module.db.session.rollback()
    return digest.hexdigest()


# ------------------------------------------------------------------ properties format
def test_valid_rows_get_proposals_and_every_bad_row_is_reported(app_module):
    rows = [HEADERS,
            [9001, "SYN-1001", 45.5, "RESIDENTIAL", None, None],        # 2  valid, new
            [9002, "SYN-1002", 0, "RESIDENTIAL", None, None],           # 3  zero area
            [9003, "SYN-1003", -4, "RESIDENTIAL", None, None],          # 4  negative
            [9004, "SYN-1004", "abc", "RESIDENTIAL", None, None],       # 5  not a number
            [9005, "SYN-1005", 12.345, "RESIDENTIAL", None, None],      # 6  3 decimals
            [9006, "SYN-1006", 30, "PARKING", None, None],              # 7  group not approved
            [9007, "SYN-1007", 30, "MEZZANINE", None, None],            # 8  group unmapped
            [None, "SYN-1008", 30, "RESIDENTIAL", None, None],          # 9  no external id
            [9009, "syn-1009 ", 30, "RESIDENTIAL", None, None],         # 10 valid (case and spaces ignored)
            [9010, "SYN 1010", 30, "2F", None, None],                   # 11 duplicate unit no...
            [9011, "syn1010", 30, "2F", None, None],                    # 12 ...ignoring case/spaces
            [9012, "SYN-1012", 30, "RESIDENTIAL", "", "#CLEAR"],        # 13 clearing a tenant
            [9012, "SYN-1013", 30, "RESIDENTIAL", None, None]]          # 14 duplicate external id with 13
    r = check(app_module, {"Properties": rows})
    assert r["format"] == "properties" and r["dryRun"] and r["databaseWrites"] == 0
    assert codes_at(r, 2) == set()
    assert "AREA_ZERO" in codes_at(r, 3) and "AREA_NEGATIVE" in codes_at(r, 4)
    assert "AREA_NOT_A_NUMBER" in codes_at(r, 5) and "AREA_PRECISION" in codes_at(r, 6)
    assert "GROUP_UNCONFIRMED" in codes_at(r, 7) and "GROUP_UNMAPPED" in codes_at(r, 8)
    assert "EXTERNAL_ID_MISSING" in codes_at(r, 9)
    assert "UNIT_NO_DUPLICATE" in codes_at(r, 11) and "UNIT_NO_DUPLICATE" in codes_at(r, 12)
    assert "PERSON_CLEAR_NOT_SUPPORTED" in codes_at(r, 13)
    assert "EXTERNAL_ID_DUPLICATE" in codes_at(r, 13) and "EXTERNAL_ID_DUPLICATE" in codes_at(r, 14)
    sheet = r["sheets"]["Properties"]
    assert sheet["rows"] == 13 and sheet["valid"] == 2      # rows 2 and 10
    assert r["proposals"]["unit"]["insert"] == 2
    assert r["readyForImport"] is False


def test_cell_states_distinguish_missing_column_blank_zero_and_clear(app_module):
    rows = [["ID", "Name", "Size", "Group", "Owner"],      # no Tenant column at all
            [9101, "SYN-2001", 0, "RESIDENTIAL", None],
            [9102, "SYN-2002", "#CLEAR", "RESIDENTIAL", ""],
            [9103, "SYN-2003", None, "RESIDENTIAL", "Synthetic Owner A"]]
    r = check(app_module, {"Properties": rows})
    states = r["cellStates"]["Properties"]
    assert states["Size"]["zero"] == 1 and states["Size"]["clear"] == 1 and states["Size"]["blank"] == 1
    assert states["Tenant"]["missing_column"] == 3
    assert states["Owner"]["blank"] == 2 and states["Owner"]["value"] == 1
    assert "VALUE_CLEAR_NOT_ALLOWED" in codes_at(r, 3, "Size")
    assert "AREA_REQUIRED" in codes_at(r, 4, "Size")          # blank is not allowed for a NEW unit


def test_people_are_never_proposed_or_merged_until_approved(app_module):
    rows = [HEADERS,
            [9201, "SYN-3001", 40, "RESIDENTIAL", "Synthetic Person One", None],
            [9202, "SYN-3002", 40, "RESIDENTIAL", "Synthetic One / Synthetic Two", "Synthetic Person Three"],
            [9203, "SYN-3003", 40, "RESIDENTIAL", "Synthetic Person One", "Synthetic Person One"],
            [9204, "SYN-3004", 40, "RESIDENTIAL", None, "Synthetic Person Four"]]
    r = check(app_module, {"Properties": rows})
    assert "PERSON_RELATIONSHIP_UNCONFIRMED" in codes_at(r, 2, "Owner")
    assert "PERSON_MULTIPLE_NAMES" in codes_at(r, 3, "Owner")
    assert "OWNER_EQUALS_TENANT" in codes_at(r, 4) and "PERSON_NAME_REPEATED" in codes_at(r, 4, "Owner")
    assert "TENANT_WITHOUT_OWNER" in codes_at(r, 5)
    assert "owner" not in r["proposals"] and "tenant" not in r["proposals"]
    assert r["readiness"]["people"]["ready"] is False
    assert r["readiness"]["properties"]["ready"] is True     # the unit data itself is fine


def test_existing_units_are_linked_by_unit_number_and_conflicts_are_detected(app_module):
    m = app_module
    with m.app.app_context():
        src = m.ImportSource(code="synthetic-legacy", name="Synthetic legacy system")
        m.db.session.add(src)
        m.db.session.flush()
        t501 = m.Unit.query.filter_by(unit_no="TEST-501").first()
        t502 = m.Unit.query.filter_by(unit_no="TEST-502").first()
        amb1 = m.Unit(unit_no="AMB X1", area_sqm=10, unit_type="STUDIO TYPE")
        amb2 = m.Unit(unit_no="AMBX1", area_sqm=10, unit_type="STUDIO TYPE")
        m.db.session.add_all([amb1, amb2])
        m.db.session.flush()
        m.db.session.add_all([m.ImportCrosswalk(source_id=src.id, entity_type="unit", external_id="7001", target_id=t501.id),
                              m.ImportCrosswalk(source_id=src.id, entity_type="unit", external_id="7002", target_id=t502.id)])
        m.db.session.commit()
        ids = (src.id, amb1.id, amb2.id)
    try:
        rows = [HEADERS,
                [7001, "TEST-501", 40, "RESIDENTIAL", None, None],     # 2 crosswalk -> same unit: update/unchanged
                [7002, "TEST-503", 40, "RESIDENTIAL", None, None],     # 3 crosswalk -> a unit with another number
                [7777, "TEST-502", 40, "RESIDENTIAL", None, None],     # 4 TEST-502 is mapped to another external id
                [7003, "TEST-504", 40, "RESIDENTIAL", None, None],     # 5 not mapped yet: link by unit number
                [7004, "AMBX 1", 40, "RESIDENTIAL", None, None]]       # 6 two units match ignoring spaces
        r = check(m, {"Properties": rows})
        assert codes_at(r, 2) == set()
        assert "CROSSWALK_UNIT_NO_MISMATCH" in codes_at(r, 3)
        assert "UNIT_MAPPED_TO_OTHER_EXTERNAL_ID" in codes_at(r, 4)
        assert codes_at(r, 5) == set()
        assert "AMBIGUOUS_UNIT_MATCH" in codes_at(r, 6)
        assert r["proposals"]["unit"]["link"] == 1
        assert r["proposals"]["unit"]["update"] + r["proposals"]["unit"]["unchanged"] == 1
        assert r["sheets"]["Properties"]["conflict"] == 3
    finally:
        with m.app.app_context():
            m.ImportCrosswalk.query.filter_by(source_id=ids[0]).delete()
            m.ImportSource.query.filter_by(id=ids[0]).delete()
            m.Unit.query.filter(m.Unit.id.in_(ids[1:])).delete(synchronize_session=False)
            m.db.session.commit()


def test_properties_layout_is_never_reported_ready_for_the_financial_migration(app_module):
    r = check(app_module, {"Properties": [HEADERS, [9301, "SYN-4001", 40, "RESIDENTIAL", None, None]]})
    fin = {b["code"] for b in r["readiness"]["financial"]["blockers"]}
    assert {"FINANCIAL_SOURCE_MISSING", "FINANCIAL_NOT_APPROVED", "FINANCIAL_CUTOFF_MISSING",
            "FINANCIAL_MODE_MISSING", "CONTROL_TOTALS_MISSING"} <= fin
    assert r["readiness"]["import"]["ready"] is False and r["readyForImport"] is False


def test_missing_required_columns_and_unconfigured_source_block(app_module):
    r = check(app_module, {"Properties": [["ID", "Name", "Group"], [1, "SYN-5001", "RESIDENTIAL"]]}, config={})
    assert any(i["code"] == "REQUIRED_COLUMN_MISSING" and i["column"] == "Size" for i in r["issues"])
    assert {p["code"] for p in r["configProblems"]} == {"SOURCE_SYSTEM_NOT_CONFIGURED"}
    blockers = {b["code"] for b in r["readiness"]["properties"]["blockers"]}
    assert {"SOURCE_SYSTEM_NOT_CONFIGURED", "UNIT_NO_UNCONFIRMED"} <= blockers
    assert "unit" not in r["proposals"]


# ------------------------------------------------------------------ CITYLAND9 export format
def export_book(m, **overrides):
    with m.app.app_context():
        t502 = m.Unit.query.filter_by(unit_no="TEST-502").first()
        bill = m.Billing.query.filter_by(unit_id=t502.id).order_by(m.Billing.billing_month).first()
        month = bill.billing_month
    sheets = {
        "Units": [["id", "unit_no", "unit_type", "area_sqm", "dues_mode", "status", "active"],
                  [1, "TEST-502", "1 BEDROOM", 50, "per_sqm", "Occupied", True],
                  [2, "SYN-EXP-1", "STUDIO TYPE", 30, "weird", "Occupied", "maybe"],
                  [3, "SYN-EXP-2", "PENTHOUSE", 0, "per_sqm", "Vacant", 1]],
        "Owners": [["id", "unit_id", "owner_name", "status"],
                   [10, 2, "Synthetic Export Owner", "Current"],
                   [11, 99, "Synthetic Orphan", "Current"],
                   [12, 2, "", "Retired"]],
        "Billing": [["id", "unit_id", "billing_month", "assessment", "penalty", "previous_balance", "amount_paid", "status"],
                    [100, 1, month, 1000, 0, 0, 0, "Unpaid"],          # 2 unit/month already in CITYLAND9
                    [101, 2, "Jan 26", 1000, 0, 0, 0, "Unpaid"],       # 3 bad month
                    [102, 2, "2026-03", "abc", 0, 0, 0, "Unpaid"],     # 4 text amount
                    [103, 2, "2026-04", -50, 0, 0, 0, "Whatever"],     # 5 negative + bad status
                    [104, 1, "2026-05", 1000, 0, 500, 300, "Unpaid"],  # 6 paid w/o payments; first bill of unit 1 has arrears
                    [105, 2, "2026-06", 1000, 25, 1200, 0, "Paid"],    # 7 penalty; carried balance; status
                    [106, 2, "2026-07", 1000, 0, 0, 1000, "Paid"]],    # 8 payment rows don't add up
        "Payments": [["id", "billing_id", "amount", "payment_date"],
                     [200, 106, 400, date(2026, 7, 10)],
                     [201, 999, 100, date(2026, 7, 10)],
                     [202, 106, 0, "not a date"]],
        "Employees": [["id", "employee_no", "full_name"], [300, "", "Synthetic Staff"]],
    }
    sheets.update(overrides)
    return sheets, month


def test_export_format_validates_values_relationships_and_financial_risks(app_module):
    m = app_module
    sheets, month = export_book(m)
    r = check(m, sheets, roundtrip=False)        # export layout as a MIGRATION from another system
    assert r["format"] == "cityland_export" and r["mode"] == "migration"
    u = lambda row, col=None: {i["code"] for i in r["issues"] if i["sheet"] == "Units" and i["row"] == row and (col is None or i["column"] == col)}
    b = lambda row: {i["code"] for i in r["issues"] if i["sheet"] == "Billing" and i["row"] == row}
    o = lambda row: {i["code"] for i in r["issues"] if i["sheet"] == "Owners" and i["row"] == row}
    p = lambda row: {i["code"] for i in r["issues"] if i["sheet"] == "Payments" and i["row"] == row}
    assert "DUES_MODE_INVALID" in u(3) and "BOOLEAN_INVALID" in u(3)
    assert "UNIT_TYPE_INVALID" in u(4) and "AREA_ZERO" in u(4)
    assert "UNRESOLVED_UNIT_REFERENCE" in o(3)
    assert {"REQUIRED_VALUE_MISSING", "STATUS_INVALID"} <= o(4)
    assert "UNIT_MONTH_EXISTS" in b(2)
    assert "MONTH_INVALID" in b(3)
    assert "NUMBER_INVALID" in b(4)
    assert {"NUMBER_NEGATIVE", "STATUS_INVALID"} <= b(5)
    assert {"PAYMENT_DETAIL_MISSING", "OPENING_ARREARS_NOT_CARRIED"} <= b(6)
    assert {"PENALTY_RECALCULATED", "PREVIOUS_BALANCE_WITH_HISTORY", "STATUS_INCONSISTENT"} <= b(7)
    assert "PAYMENT_TOTAL_MISMATCH" in b(8)
    assert "UNRESOLVED_BILL_REFERENCE" in p(3)
    assert {"AMOUNT_NOT_POSITIVE", "DATE_INVALID"} <= p(4)
    assert any(i["sheet"] == "Employees" and i["code"] == "EMPLOYEE_NO_MISSING" for i in r["issues"])
    fin = {x["code"] for x in r["readiness"]["financial"]["blockers"]}
    assert {"ADVANCE_LEDGER_MISSING", "RECEIPTS_NOT_IN_SOURCE", "FINANCIAL_NOT_APPROVED"} <= fin
    assert r["readyForImport"] is False


def test_closed_periods_issued_and_hand_corrected_bills_are_blocked(app_module):
    m = app_module
    with m.app.app_context():
        t503 = m.Unit.query.filter_by(unit_no="TEST-503").first()
        bill = m.Billing.query.filter_by(unit_id=t503.id).order_by(m.Billing.billing_month.desc()).first()
        before = (bill.soa_manual_override, bill.issued_at)
        bill.soa_manual_override, bill.issued_at = True, datetime(2026, 1, 1)
        setting = m.Setting.query.filter_by(key="books_closed_through").first() or m.Setting(key="books_closed_through")
        old_closed = setting.value
        setting.value = bill.billing_month
        m.db.session.add(setting)
        m.db.session.commit()
        bid, month = bill.id, bill.billing_month
    try:
        sheets = {"Units": [["id", "unit_no"], [1, "TEST-503"]],
                  "Billing": [["id", "unit_id", "billing_month", "assessment"], [50, 1, month, 10]],
                  "WaterReadings": [["id", "unit_id", "reading_month", "previous_reading", "current_reading"], [60, 1, month, 30, 20]]}
        r = check(m, sheets, roundtrip=False)
        bill_codes = {i["code"] for i in r["issues"] if i["sheet"] == "Billing"}
        assert {"PERIOD_CLOSED", "BILL_ISSUED", "BILL_CORRECTED_BY_HAND", "UNIT_MONTH_EXISTS"} <= bill_codes
        water_codes = {i["code"] for i in r["issues"] if i["sheet"] == "WaterReadings"}
        assert {"PERIOD_CLOSED", "READING_DECREASES"} <= water_codes
    finally:
        with m.app.app_context():
            bill = m.db.session.get(m.Billing, bid)
            bill.soa_manual_override, bill.issued_at = before
            m.Setting.query.filter_by(key="books_closed_through").first().value = old_closed
            m.db.session.commit()


def test_financial_readiness_needs_every_approved_decision_and_matching_control_totals(app_module):
    sheets = {"Units": [["id", "unit_no", "area_sqm", "unit_type"], [1, "SYN-FIN-1", 40, "STUDIO TYPE"]],
              "Billing": [["id", "unit_id", "billing_month", "assessment", "amount_paid", "status"], [1, 1, "2026-01", 1000, 0, "Unpaid"]],
              "Payments": [["id", "billing_id", "amount", "payment_date"]]}
    approved = {**CONFIG, "financial": {"approved": True, "approved_by": "synthetic", "approved_on": "2026-10-08",
                                        "cutoff_date": "2026-01-31", "mode": "full_history",
                                        "scope": {"first_month": "2026-01", "last_month": "2026-01", "entities": ["billing"]},
                                        "control_totals": {"billing_rows": 1, "billing_charges_total": "1000",
                                                           "payment_rows": 0, "payment_total": "0"}}}
    r = check(app_module, sheets, config=approved, roundtrip=False)
    fin = {x["code"] for x in r["readiness"]["financial"]["blockers"]}
    assert fin == {"ADVANCE_LEDGER_MISSING", "RECEIPTS_NOT_IN_SOURCE"}    # never invented, so still blocked
    wrong = {**approved, "financial": {**approved["financial"], "control_totals": {**approved["financial"]["control_totals"], "billing_charges_total": "999"}}}
    r = check(app_module, sheets, config=wrong, roundtrip=False)
    assert "CONTROL_TOTALS_MISMATCH" in {x["code"] for x in r["readiness"]["financial"]["blockers"]}
    opening = {**approved, "financial": {**approved["financial"], "mode": "opening_balances", "cutoff_date": "2026-03-31"}}
    r = check(app_module, sheets, config=opening, roundtrip=False)
    assert "OPENING_BALANCE_UNSUPPORTED" in {x["code"] for x in r["readiness"]["financial"]["blockers"]}
    assert any(i["code"] == "HISTORY_IN_OPENING_BALANCE_MODE" for i in r["issues"])


# ------------------------------------------------------------------ no writes, privacy, access
def test_check_leaves_every_table_unchanged_and_refuses_writes(app_module):
    m = app_module
    sheets, _ = export_book(m)
    before = snapshot(m)
    check(m, sheets)
    check(m, sheets, roundtrip=False)
    check(m, {"Properties": [HEADERS, [1, "SYN-6001", 40, "RESIDENTIAL", "Synthetic Person", None]]})
    assert snapshot(m) == before
    with m.app.app_context():
        with pytest.raises(ic.DryRunWriteError):
            with ic.no_writes(m.db):
                m.db.session.add(m.Setting(key="synthetic_write_attempt", value="x"))
                m.db.session.flush()
        with pytest.raises(ic.DryRunWriteError):
            with ic.no_writes(m.db):
                m.db.session.execute(text("DELETE FROM setting WHERE key = 'nothing'"))
        assert m.Setting.query.filter_by(key="synthetic_write_attempt").count() == 0


def test_check_only_endpoint_is_superadmin_only_csrf_protected_and_writes_nothing(app_module, superadmin, tmp_path, monkeypatch):
    monkeypatch.setenv("IMPORT_CONFIG_PATH", str(tmp_path / "none.json"))
    book = lambda: (xlsx({"Properties": [HEADERS, [8801, "SYN-7001", 40, "RESIDENTIAL", "Synthetic Secret Name", None]]}), "client.xlsx")
    # no CSRF token
    resp = superadmin.post("/api/admin/system/import/check", data={"file": book()}, content_type="multipart/form-data")
    assert resp.status_code == 403
    client = api(superadmin)
    before = snapshot(app_module)
    with app_module.app.app_context():
        backups_before = app_module.AuditLog.query.count()
    resp = client.post("/api/admin/system/import/check", data={"file": book()}, content_type="multipart/form-data")
    assert resp.status_code == 200, resp.get_data(as_text=True)
    body = resp.get_data(as_text=True)
    assert resp.headers["Cache-Control"] == "no-store"
    assert "Synthetic Secret Name" not in body and "SYN-7001" not in body and "8801" not in body
    report = resp.get_json()
    assert report["dryRun"] is True and report["configLoaded"] is False
    assert snapshot(app_module) == before
    with app_module.app.app_context():
        assert app_module.AuditLog.query.count() == backups_before
    resp = client.post("/api/admin/system/import/check", data={"file": (io.BytesIO(b"not excel"), "x.xlsx")}, content_type="multipart/form-data")
    assert resp.status_code == 400
    staff, _ = login(app_module, "test_staff", PASSWORD)
    staff = api(staff)
    resp = staff.post("/api/admin/system/import/check", data={"file": book()}, content_type="multipart/form-data")
    assert resp.status_code == 403


def test_children_of_an_invalid_unit_or_bill_are_not_proposed(app_module):
    sheets = {"Units": [["id", "unit_no", "area_sqm", "unit_type"], [1, "SYN-PAR-1", 0, "STUDIO TYPE"],
                        [2, "SYN-PAR-2", 30, "STUDIO TYPE"]],
              "Owners": [["id", "unit_id", "owner_name"], [10, 1, "Synthetic Child Owner"]],
              "Billing": [["id", "unit_id", "billing_month", "assessment", "amount_paid", "status"],
                          [20, 2, "2026-13", 100, 100, "Paid"]],
              "Payments": [["id", "billing_id", "amount", "payment_date"], [30, 20, 100, date(2026, 1, 5)]]}
    r = check(app_module, sheets, roundtrip=False)
    assert {i["code"] for i in r["issues"] if i["sheet"] == "Owners"} == {"PARENT_RECORD_NOT_VALID"}
    assert "PARENT_RECORD_NOT_VALID" in {i["code"] for i in r["issues"] if i["sheet"] == "Payments"}
    assert "owner" not in r["proposals"] and "payment" not in r["proposals"]


# ------------------------------------------------------------------ the lock on the classic import
def own_export(superadmin):
    """CITYLAND9's own export of this (test) database, as the Settings download produces it."""
    from openpyxl import load_workbook
    resp = superadmin.get("/database/export.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.data))
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def rows_of(sheets, sheet):
    return sheets[sheet][0], sheets[sheet][1:]


def test_an_unchanged_export_of_this_database_passes_the_gate_and_imports(app_module, superadmin):
    sheets = own_export(superadmin)
    r = check(app_module, sheets)
    assert r["mode"] == "cityland_roundtrip"
    assert r["importGate"] == {"allowed": True, "blockers": []}, r["issueCounts"]
    client = api(superadmin)
    resp = client.post("/api/admin/system/import", data={"file": (xlsx(sheets), "export.xlsx")}, content_type="multipart/form-data")
    assert resp.status_code == 200, resp.get_data(as_text=True)
    assert all(c["inserted"] == 0 for c in resp.get_json()["counts"].values())


def test_the_gate_refuses_every_unsafe_change_and_the_import_writes_nothing(app_module, superadmin):
    m = app_module
    with m.app.app_context():
        unit = m.Unit.query.filter_by(unit_no="TEST-504").first()
        original_dues = unit.manual_monthly_dues
        unit.manual_monthly_dues = 555
        m.db.session.commit()
        manual_id = unit.id
    try:
        _unsafe_changes_are_refused(m, superadmin, manual_id)
    finally:
        with m.app.app_context():
            m.db.session.get(m.Unit, manual_id).manual_monthly_dues = original_dues
            m.db.session.commit()


def _unsafe_changes_are_refused(m, superadmin, manual_id):
    sheets = own_export(superadmin)
    h, units = rows_of(sheets, "Units")
    col = {c: i for i, c in enumerate(h)}
    first_other = next(r for r in units if r[col["id"]] != manual_id)
    first_other[col["unit_no"]] = "RENAMED-BY-ID"                    # the id now names another unit number
    for row in units:
        if row[col["id"]] == manual_id:
            row[col["manual_monthly_dues"]] = None                   # blank: keeps the saved amount (not a refusal)
    sheets["Units"].append([999999, "NEW-UNIT-ROW"] + [None] * (len(h) - 2))   # a record this database doesn't have
    oh, owners = rows_of(sheets, "Owners")
    ocol = {c: i for i, c in enumerate(oh)}
    untouched = [r[col["id"]] for r in units if r is not first_other and r[col["id"]] not in (manual_id, owners[0][ocol["unit_id"]])
                 and (r[col["unit_type"]] or "") not in ("PARKING", "STORAGE")]
    owners[0][ocol["unit_id"]] = untouched[0]                    # the owner id would be moved to another (valid) unit
    r = check(m, sheets)
    codes = {b["code"] for b in r["importGate"]["blockers"]}
    assert r["importGate"]["allowed"] is False
    assert {"CROSSWALK_UNIT_NO_MISMATCH", "NEW_RECORD_NOT_ALLOWED", "ID_POINTS_TO_OTHER_RECORD"} <= codes, codes
    client = api(superadmin)
    before = snapshot(m)
    resp = client.post("/api/admin/system/import", data={"file": (xlsx(sheets), "export.xlsx")}, content_type="multipart/form-data")
    assert resp.status_code == 409 and "nothing was changed" in resp.get_json()["error"]["message"]
    assert snapshot(m) == before


def test_closed_months_and_issued_bills_lock_only_when_the_import_would_change_them(app_module, superadmin):
    m = app_module
    sheets = own_export(superadmin)
    h, bills = rows_of(sheets, "Billing")
    col = {c: i for i, c in enumerate(h)}
    bid = bills[0][col["id"]]
    with m.app.app_context():
        bill = m.db.session.get(m.Billing, bid)
        old_issued = bill.issued_at
        bill.issued_at = datetime(2026, 1, 1)
        setting = m.Setting.query.filter_by(key="books_closed_through").first() or m.Setting(key="books_closed_through")
        old_closed = setting.value
        setting.value = bill.billing_month
        m.db.session.add(setting)
        m.db.session.commit()
    try:
        r = check(m, sheets)
        assert r["importGate"]["allowed"] is True, r["importGate"]   # unchanged: the import would not alter it
        bills[0][col["assessment"]] = 12345
        r = check(m, sheets)
        codes = {b["code"] for b in r["importGate"]["blockers"]}
        assert {"PERIOD_CLOSED", "BILL_ISSUED"} <= codes, codes
    finally:
        with m.app.app_context():
            bill = m.db.session.get(m.Billing, bid)
            bill.issued_at = old_issued
            m.Setting.query.filter_by(key="books_closed_through").first().value = old_closed
            m.db.session.commit()


def test_another_layout_is_never_importable(app_module, superadmin):
    client = api(superadmin)
    rows = [HEADERS, [1, "SYN-LOCK-1", 40, "RESIDENTIAL", None, None]]
    before = snapshot(app_module)
    resp = client.post("/api/admin/system/import", data={"file": (xlsx({"Properties": rows}), "client.xlsx")}, content_type="multipart/form-data")
    assert resp.status_code == 409 and "Only CITYLAND9's own export" in resp.get_json()["error"]["message"]
    assert snapshot(app_module) == before
    r = check(app_module, {"Properties": rows})
    assert r["importGate"]["allowed"] is False and r["readyForImport"] is False


def test_round_trip_reports_unchanged_rows_and_blocks_hidden_changes(app_module, superadmin):
    sheets = own_export(superadmin)
    r = check(app_module, sheets)
    assert r["proposals"]["unit"]["unchanged"] == len(sheets["Units"]) - 1 and r["proposals"]["unit"]["update"] == 0
    assert r["proposals"]["owner"]["unchanged"] == len(sheets["Owners"]) - 1
    h, units = rows_of(sheets, "Units")
    col = {c: i for i, c in enumerate(h)}
    no_owner = next(row for row in units if row[col["unit_type"]] == "PARKING")
    no_owner[col["owner_name"]] = "Synthetic New Owner"           # would create a new owner record
    r = check(app_module, sheets)
    issues = {(i["column"], i["code"]) for i in r["issues"]}
    assert ("owner_name", "NEW_RECORD_NOT_ALLOWED") in issues
    assert r["importGate"]["allowed"] is False



def test_blank_cells_keep_saved_values_in_a_real_import(app_module, superadmin):
    """The classic import used to turn blanks into 0 / today / unpaid / active and to clear parking
    assignments. Blank now means "keep what is saved"; checked through an actual import."""
    m = app_module
    with m.app.app_context():
        unit = m.Unit.query.filter(m.Unit.assigned_parking_unit_id.isnot(None)).first()
        reading = m.WaterReading.query.order_by(m.WaterReading.id).first()
        old = {"dues": unit.manual_monthly_dues, "active": unit.active, "reading": (reading.reading_date, reading.paid, reading.rate)}
        unit.manual_monthly_dues, unit.active = 777, False
        reading.reading_date, reading.paid = date(2026, 2, 3), True
        m.db.session.commit()
        uid, parking, rid = unit.id, unit.assigned_parking_unit_id, reading.id
        prev, cur = reading.previous_reading, reading.current_reading
    try:
        sheets = own_export(superadmin)
        h, units = rows_of(sheets, "Units")
        col = {c: i for i, c in enumerate(h)}
        row = next(r for r in units if r[col["id"]] == uid)
        for c in ("manual_monthly_dues", "active", "assigned_parking_unit_id", "floor", "area_sqm"):
            row[col[c]] = None
        wh, readings = rows_of(sheets, "WaterReadings")
        wcol = {c: i for i, c in enumerate(wh)}
        wrow = next(r for r in readings if r[wcol["id"]] == rid)
        for c in ("previous_reading", "current_reading", "rate", "reading_date", "paid"):
            wrow[wcol[c]] = None
        r = check(m, sheets)
        assert r["importGate"]["allowed"] is True, r["importGate"]
        resp = api(superadmin).post("/api/admin/system/import", data={"file": (xlsx(sheets), "export.xlsx")}, content_type="multipart/form-data")
        assert resp.status_code == 200, resp.get_data(as_text=True)
        with m.app.app_context():
            u = m.db.session.get(m.Unit, uid)
            w = m.db.session.get(m.WaterReading, rid)
            assert float(u.manual_monthly_dues) == 777 and u.active is False and u.assigned_parking_unit_id == parking
            assert (w.reading_date, w.paid, w.previous_reading, w.current_reading) == (date(2026, 2, 3), True, prev, cur)
    finally:
        with m.app.app_context():
            u = m.db.session.get(m.Unit, uid)
            u.manual_monthly_dues, u.active = old["dues"], old["active"]
            w = m.db.session.get(m.WaterReading, rid)
            w.reading_date, w.paid, w.rate = old["reading"]
            m.db.session.commit()


def test_the_importer_itself_never_uses_a_file_id_as_a_primary_key_and_errors_are_plain(app_module):
    """Defence in depth below the gate: run_excel_import directly (the API always gates first)."""
    m = app_module
    with m.app.app_context():
        top = m.db.session.query(m.func.max(m.Unit.id)).scalar()
    foreign_id = top + 5000
    book = xlsx({"Units": [["id", "unit_no", "area_sqm", "unit_type"], [foreign_id, "SYN-PK-1", 30, "STUDIO TYPE"]]})
    from werkzeug.datastructures import FileStorage
    with m.app.test_request_context():
        ok, _, _ = m.run_excel_import(FileStorage(stream=book, filename="pk.xlsx"))
    with m.app.app_context():
        created = m.Unit.query.filter_by(unit_no="SYN-PK-1").first()
        try:
            assert ok and created is not None and created.id != foreign_id
        finally:
            m.db.session.delete(created)
            m.db.session.commit()
    with m.app.app_context():
        t502 = m.Unit.query.filter_by(unit_no="TEST-502").first()
        bill = m.Billing.query.filter_by(unit_id=t502.id).first()
        bid, month, uid = bill.id, bill.billing_month, t502.id
        other = m.Unit.query.filter(m.Unit.id != uid, m.Unit.unit_type.notin_(["PARKING", "STORAGE"])).first().id
    for sheets, expected in (
        ({"Units": [["id", "unit_no"], [uid, "TEST-502"], [other, "OTHER"]],
          "Billing": [["id", "unit_id", "billing_month", "assessment"], [bid, other, month, 1]]}, "belongs to another unit or month"),
        ({"Units": [["id", "unit_no"], [uid, "TEST-502"]],
          "Billing": [["id", "unit_id", "billing_month", "assessment"], [999999, uid, month, 1]]}, "already has a bill for"),
        ({"Units": [["id", "unit_no"], [uid, "TEST-502"]],
          "Billing": [["id", "unit_id", "billing_month", "assessment"], [None, uid, "2026-13", 1]]}, "could not be saved"),
    ):
        before = snapshot(m)
        with m.app.test_request_context():
            ok, message, _ = m.run_excel_import(FileStorage(stream=xlsx(sheets), filename="x.xlsx"))
        assert not ok and expected in message and "Traceback" not in message and "sqlite3" not in message and "IntegrityError" not in message, message
        tables_before_without_backups = before
        assert snapshot(m) == tables_before_without_backups
