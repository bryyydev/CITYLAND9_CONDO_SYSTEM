"""Property / Financial Reports (/api/reports): permissions, periods, the corrected totals (R1-R3),
voided and receipt-less payments, expenses, the Excel export and the retired classic page."""
import secrets
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest

from conftest import PASSWORD, SA_PASSWORD, SA_USERNAME, api, login

D = Decimal
FROM, TO = "2024-03-01", "2024-04-30"     # months no other test uses


@pytest.fixture
def sa(app_module):
    return api(login(app_module, SA_USERNAME, SA_PASSWORD)[0])


def report(client, **params):
    q = "&".join(f"{k}={v}" for k, v in {"period": "custom", "from": FROM, "to": TO, **params}.items())
    resp = client.get(f"/api/reports?{q}")
    assert resp.status_code == 200, resp.get_json()
    return resp.get_json()


@pytest.fixture
def world(app_module):
    """A unit with two unpaid months (3,000 each, bills dated in the report period), a partly paid water
    reading paid in two months, a voided receipt, a payment without receipt and an expense. Other
    units are inactive while it is built. Everything is removed afterwards."""
    m = app_module
    tok = lambda: secrets.token_urlsafe(24)  # noqa: E731
    with m.app.test_request_context():
        u = m.Unit(unit_no="RPT-1", unit_type="1 BEDROOM", area_sqm=40, active=True, status="Occupied")
        m.db.session.add(u); m.db.session.commit()
        others = [x.id for x in m.Unit.query.filter(m.Unit.id != u.id, m.Unit.active.is_(True)).all()]
        m.Unit.query.filter(m.Unit.id.in_(others)).update({"active": False}, synchronize_session=False)
        saved = {}
        for key, value in (("penalty_rate", "10"), ("penalty_include_condo", "1")):
            row = m.Setting.query.filter_by(key=key).first()
            saved[key] = row.value if row else None
            if row:
                row.value = value
            else:
                m.db.session.add(m.Setting(key=key, value=value))
        m.db.session.commit()
        for month in ("2024-03", "2024-04"):
            m.generate_bills(month)
        bills = {b.billing_month: b for b in m.Billing.query.filter_by(unit_id=u.id).all()}
        bills["2024-03"].created_at = m.datetime(2024, 3, 5, 3, 0)
        bills["2024-04"].created_at = m.datetime(2024, 4, 5, 3, 0)
        m.db.session.commit()
        reading, _, _ = m.save_water_reading(unit_id=u.id, month="2024-02", previous=D("0"), current=D("10"), rate=D("50"),
                                             reading_date=date(2024, 2, 28))                                   # 500.00
        m.record_water_payment(reading.id, amount=D("200"), payment_method="CASH", payment_type="PARTIAL", reference="",
                               paid_date=date(2024, 3, 10), form_token=tok())
        m.record_water_payment(reading.id, amount=D("300"), payment_method="ONLINE", payment_type="FULL", reference="GC-1",
                               paid_date=date(2024, 4, 10), form_token=tok())
        bad = m.record_bill_payment(bills["2024-03"].id, amount=D("1000"), payment_method="CHECK", payment_type="PARTIAL", reference="CHK-9",
                                    remarks="", payment_date=date(2024, 3, 20), form_token=tok())["receipt"]
        m.void_receipt_record(bad.id, "Bounced check during report test", tok())
        legacy_pay = m.Payment(billing_id=bills["2024-04"].id, amount=D("700"), payment_date=date(2024, 4, 15), payment_method="CASH")
        m.db.session.add(legacy_pay)
        bills["2024-04"].amount_paid = D("700")
        exp = m.Expense(expense_date=date(2024, 4, 2), category="Repairs", description="Report test pump", amount=D("1200"))
        m.db.session.add(exp)
        m.db.session.commit()
        ids = {"unit": u.id, "reading": reading.id, "expense": exp.id, "others": others, "settings": saved}
    yield ids
    with m.app.app_context():
        for r in m.Receipt.query.filter_by(unit_id=ids["unit"]).all():
            m.db.session.delete(r)
        m.db.session.flush()
        bill_ids = [b.id for b in m.Billing.query.filter_by(unit_id=ids["unit"]).all()]
        m.Payment.query.filter(m.Payment.billing_id.in_(bill_ids)).delete(synchronize_session=False)
        m.AdvancePayment.query.filter_by(unit_id=ids["unit"]).delete()
        m.Billing.query.filter_by(unit_id=ids["unit"]).delete()
        m.WaterReading.query.filter_by(id=ids["reading"]).delete()
        m.Expense.query.filter_by(id=ids["expense"]).delete()
        m.Unit.query.filter_by(id=ids["unit"]).delete()
        m.Unit.query.filter(m.Unit.id.in_(ids["others"])).update({"active": True}, synchronize_session=False)
        for key, value in ids["settings"].items():
            if value is None:
                m.Setting.query.filter_by(key=key).delete()
            else:
                m.Setting.query.filter_by(key=key).first().value = value
        m.db.session.commit()


@pytest.mark.parametrize("role,status", [("admin", 200), ("accounting", 200), ("manager", 403), ("staff", 403), ("resident", 403)])
def test_permissions(app_module, role, status):
    client = api(login(app_module, f"test_{role}", PASSWORD)[0])
    assert client.get("/api/reports").status_code == status
    assert client.get("/api/reports/export.xlsx").status_code == (200 if status == 200 else 403)


def test_signed_out_is_refused(app_module):
    assert app_module.app.test_client().get("/api/reports").status_code == 401


def test_period_presets_and_bad_dates(sa):
    today = date.today()
    monthly = sa.get("/api/reports").get_json()
    assert monthly["period"] == "monthly" and monthly["from"] == today.replace(day=1).isoformat()
    daily = sa.get("/api/reports?period=daily").get_json()
    assert daily["from"] == daily["to"] == today.isoformat()
    weekly = sa.get("/api/reports?period=weekly").get_json()
    assert date.fromisoformat(weekly["from"]).weekday() == 0 and (date.fromisoformat(weekly["to"]) - date.fromisoformat(weekly["from"])).days == 6
    yearly = sa.get("/api/reports?period=yearly").get_json()
    assert (yearly["from"], yearly["to"]) == (f"{today.year}-01-01", f"{today.year}-12-31")
    swapped = sa.get("/api/reports?period=custom&from=2024-04-30&to=2024-03-01").get_json()
    assert (swapped["from"], swapped["to"]) == ("2024-03-01", "2024-04-30")
    assert sa.get("/api/reports?period=custom&from=03/01/2024").status_code == 400
    assert sa.get("/api/reports?period=custom&from=1990-01-01&to=2024-01-01").status_code == 400


def owed_elsewhere(m, unit_id):
    """What the OTHER units still owe (their bills share the test database)."""
    with m.app.test_request_context():
        m.reset_financial_caches()
        owed = {}
        for b in m.Billing.query.filter(m.Billing.unit_id != unit_id).all():
            if m.bill_unpaid_part(b) > 0:
                owed[b.unit_id] = owed.get(b.unit_id, D("0")) + m.bill_unpaid_part(b)
        return sum(owed.values(), D("0")), len(owed)


def test_totals_are_counted_once(app_module, sa, world):
    r = report(sa)
    # R1: charges of the two bills (3,000 dues each + April's 10% penalty on March's unpaid dues),
    # NOT also April's previous balance (the old page: 9,300).
    assert r["billing"]["bills"] == 2 and r["billing"]["charged"] == "6300.00"
    assert r["billing"]["charges"]["condo"] == "6000.00" and r["billing"]["charges"]["penalty"] == "300.00"
    # R3: the water reading was paid 200 in March and 300 in April: each counted in its own month.
    march = report(sa, to="2024-03-31")
    assert march["collections"]["water"] == "200.00"
    april = report(sa, **{"from": "2024-04-01"})
    assert april["collections"]["water"] == "300.00"
    # Voided receipts are not money received; they are reported on their own.
    assert r["collections"]["voided"] == {"count": 1, "amount": "1000.00"}
    # A payment recorded without a receipt is still counted, and flagged.
    assert r["collections"]["withoutReceipt"] == 1 and r["collections"]["bills"] == "700.00"
    assert r["collections"]["total"] == "1200.00"
    assert r["collections"]["byMethod"] == {"CASH": "900.00", "CHECK": "0.00", "ONLINE": "300.00"}
    no_receipt = [t for t in r["transactions"] if t["receiptNo"] is None]
    assert len(no_receipt) == 1 and no_receipt[0]["amount"] == "700.00" and no_receipt[0]["unitNo"] == "RPT-1"
    assert all(t["receiptNo"].startswith("OR-") for t in r["transactions"] if t["receiptNo"])
    # Expenses and net cash flow.
    assert r["expenses"]["total"] == "1200.00" and r["expenses"]["byCategory"] == {"Repairs": "1200.00"}
    assert r["netCashFlow"] == "0.00"
    # R2: the unit owes 6,300 - 700 = 5,600 (counted once; the old page added March's 3,000 again).
    other_total, other_units = owed_elsewhere(app_module, world["unit"])
    assert D(r["receivables"]["total"]) - other_total == D("5600.00") and r["receivables"]["units"] == other_units + 1
    assert r["property"]["residentialUnits"] == 1 and r["property"]["occupied"] == 1


def test_report_changes_nothing(app_module, sa, world):
    m = app_module
    with m.app.app_context():
        before = (m.Billing.query.count(), m.Payment.query.count(), m.Receipt.query.count(), m.AuditLog.query.count())
    report(sa)
    with m.app.app_context():
        assert (m.Billing.query.count(), m.Payment.query.count(), m.Receipt.query.count(), m.AuditLog.query.count()) == before


def test_excel_export(app_module, sa, world):
    from openpyxl import load_workbook
    resp = sa.get(f"/api/reports/export.xlsx?period=custom&from={FROM}&to={TO}")
    assert resp.status_code == 200 and resp.data[:2] == b"PK"
    wb = load_workbook(BytesIO(resp.data))
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows() if row[0].value}
    assert summary["Charges billed"] == 6300 and summary["Collections (total)"] == 1200
    assert D(str(summary["Receivables now"])) - owed_elsewhere(app_module, world["unit"])[0] == D("5600")
    tx = list(wb["Transactions"].iter_rows(values_only=True))
    assert tx[0][2] == "OR No." and any(r[2] == "(no receipt)" for r in tx[1:]) and any(r[0] == "Expense" and r[7] == -1200 for r in tx[1:])
    with app_module.app.app_context():
        assert app_module.AuditLog.query.filter(app_module.AuditLog.action.like("Exported the condo report 2024-03-01%")).count() >= 1


def test_classic_page_redirects_by_role(app_module):
    admin = login(app_module, "test_admin", PASSWORD)[0]
    assert admin.get("/reports?period=custom&start_date=2024-03-01&end_date=2024-04-30").headers["Location"] == \
        "/app/admin/reports?period=custom&from=2024-03-01&to=2024-04-30"
    acc = login(app_module, "test_accounting", PASSWORD)[0]
    assert acc.get("/reports").headers["Location"] == "/app/accounting/financial-reports"
    loc = acc.get("/reports/export.xlsx?start_date=2024-03-01&end_date=2024-04-30").headers["Location"]
    assert loc == "/api/reports/export.xlsx?period=custom&from=2024-03-01&to=2024-04-30"
    staff = login(app_module, "test_staff", PASSWORD)[0]
    assert staff.get("/reports").status_code == 302 and "/app/" not in staff.get("/reports").headers["Location"].split("?")[0][:5]


def test_monthly_trend(app_module, sa, world):
    assert sa.get("/api/reports/monthly?months=0").status_code == 400
    assert sa.get("/api/reports/monthly?months=abc").status_code == 400
    assert sa.get("/api/reports/monthly?months=37").status_code == 400
    six = sa.get("/api/reports/monthly").get_json()["months"]
    assert len(six) == 6 and six[-1]["month"] == date.today().strftime("%Y-%m")
    rows = {r["month"]: r for r in sa.get("/api/reports/monthly?months=36").get_json()["months"]}
    mar, apr = rows["2024-03"], rows["2024-04"]
    # Billed: the month's own charges (April: dues + 10% penalty on March), no earlier balances.
    assert (mar["billed"], apr["billed"]) == ("3000.00", "3300.00")
    # Collected: March 200 water (the bounced 1,000 check is voided); April 300 water + 700 without receipt.
    assert (mar["collected"], apr["collected"]) == ("200.00", "1000.00")
    assert abs(apr["collectionRate"] - 1000 / 3300) < 1e-9
    # Receivables at month end for this unit: March 3,000 (only bills up to March; the voided payment
    # is not counted), April 6,300 - 700.
    # (the other test units' bills are from 2026, so in 2024 only this unit owed anything)
    assert (mar["outstanding"], apr["outstanding"]) == ("3000.00", "5600.00")


@pytest.mark.parametrize("role,status", [("accounting", 200), ("admin", 200), ("staff", 403), ("manager", 403)])
def test_monthly_permissions(app_module, role, status):
    assert api(login(app_module, f"test_{role}", PASSWORD)[0]).get("/api/reports/monthly").status_code == status
