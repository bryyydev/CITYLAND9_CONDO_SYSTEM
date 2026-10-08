"""Dry-run ("Check only") validation of a data-migration workbook. It never writes.

Guide: CITYLAND9_IMPORT_SAFEGUARDS.md. Used by:
  * POST /api/admin/system/import/check  (Superadmin, CSRF-protected, same upload limits as the import)
  * database/import_check.py             (command line, for a staging copy on the server PC)

What "no writes" means here: no INSERT/UPDATE/DELETE, no backup, no audit entry, no account and no
file is created. A guard (no_writes) refuses any write statement on the database connection while a
check runs, the ORM session refuses to flush, and the session is rolled back afterwards.

Privacy: the report names a sheet, row, column and error code only. It never contains names,
contacts, unit numbers, ids or amounts from the workbook. Source values are read in memory and
discarded.

Two workbook formats are recognised:
  properties       the client's property list: one sheet (default "Properties") with an external id,
                   unit number, size (sqm), group, owner and tenant columns (import_config.json maps them)
  cityland_export  CITYLAND9's own Excel export (Units, Owners, Tenants, Billing, Payments,
                   WaterReadings, Employees, Expenses)

Modes: a CITYLAND9 export layout is checked by default as a ROUND TRIP (cityland_roundtrip): an edit of
THIS database's own export, its ids being this database's ids. The report's "importGate" says whether the
classic import (Settings > Import workbook, run_excel_import) may run on the file: only when every row
updates an existing record whose id matches, nothing new is added, no blank cell would reset a value, and
no closed month or issued/hand-corrected bill would change. Every other layout is a MIGRATION (check only).

Row results: valid (a confirmed matching rule gives an insert/update/link/unchanged proposal),
blocked (invalid data), conflict (contradicts the database or the crosswalk), unresolved (needs a
decision that has not been approved). Every supported row is checked; nothing stops at the first error.
"""
import hashlib
import json
import os
import re
import threading
import zipfile
from contextlib import contextmanager
from datetime import date, datetime
from decimal import Decimal

from ..core.numbers import InputError, parse_decimal

CLEAR_MARKER = "#CLEAR"
RANK = {"blocked": 3, "conflict": 2, "unresolved": 1, "warning": 0}
MAX_ISSUES_LISTED = 1000
UNIT_TYPES = ("STUDIO TYPE", "1 BEDROOM", "2 BEDROOM", "3 BEDROOM", "PARKING", "STORAGE")
DUES_MODES = ("per_sqm", "manual")
UNIT_STATUSES = ("Occupied", "Vacant")
PERSON_STATUSES = ("Current", "Past")
BILL_STATUSES = ("Unpaid", "Partially Paid", "Paid", "Overdue")   # what bill_status() stores
PAYMENT_METHODS = ("CASH", "CHECK", "ONLINE")
FINANCIAL_MODES = ("full_history", "opening_balances")
MULTI_PERSON = re.compile(r"[/&;,+\n]|\band\b", re.IGNORECASE)
MONTH_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")

# Every code with its fixed, value-free message.
MESSAGES = {
    # file / format / configuration
    "FILE_NOT_XLSX": "That file is not a valid Excel .xlsx workbook.",
    "FILE_TOO_LARGE": "The workbook is too large when unpacked.",
    "SHEET_TOO_MANY_ROWS": "The sheet has more rows than the import limit.",
    "SHEET_TOO_MANY_COLUMNS": "The sheet has more columns than the import limit.",
    "FORMAT_UNKNOWN": "No supported sheet layout was found (a Properties sheet or CITYLAND9's export sheets).",
    "REQUIRED_COLUMN_MISSING": "A required column is missing from the sheet.",
    "REQUIRED_SHEET_MISSING": "A sheet required for this format is missing.",
    "SOURCE_SYSTEM_NOT_CONFIGURED": "No stable source-system code is configured, so external ids have no namespace.",
    "CROSSWALK_UNAVAILABLE": "The import crosswalk table is missing (database schema older than 0012).",
    # identity
    "EXTERNAL_ID_MISSING": "The external id is blank.",
    "EXTERNAL_ID_FORMAT": "The external id is not a whole number or text.",
    "EXTERNAL_ID_DUPLICATE": "The same external id appears on more than one row of this sheet.",
    "UNIT_NO_MISSING": "The unit number is blank.",
    "UNIT_NO_TOO_LONG": "The unit number is longer than 40 characters.",
    "UNIT_NO_DUPLICATE": "The same unit number (ignoring case and spaces) appears on more than one row.",
    "UNIT_NO_UNCONFIRMED": "It is not confirmed that this column is the unit number shown on SOAs.",
    "CROSSWALK_TARGET_MISSING": "The crosswalk points to a CITYLAND9 record that no longer exists.",
    "CROSSWALK_UNIT_NO_MISMATCH": "The crosswalk maps this external id to a unit with a different unit number.",
    "UNIT_MAPPED_TO_OTHER_EXTERNAL_ID": "The existing unit with this unit number is already mapped to a different external id.",
    "AMBIGUOUS_UNIT_MATCH": "More than one existing unit matches this unit number when case and spaces are ignored.",
    # values
    "VALUE_CLEAR_NOT_ALLOWED": "This field is required and can't be cleared.",
    "AREA_REQUIRED": "Size (sqm) is required for a new unit.",
    "AREA_NOT_A_NUMBER": "Size (sqm) is not a finite number.",
    "AREA_PRECISION": "Size (sqm) has more than 2 decimal places.",
    "AREA_ZERO": "Size (sqm) is 0; a unit with no area would be billed nothing.",
    "AREA_NEGATIVE": "Size (sqm) is negative.",
    "GROUP_MISSING": "The group is blank.",
    "GROUP_UNMAPPED": "The group label has no approved meaning in the configuration.",
    "GROUP_UNCONFIRMED": "The group label's meaning is not approved yet.",
    "UNIT_TYPE_INVALID": "The unit type is not one of CITYLAND9's types.",
    "DUES_MODE_INVALID": "The dues mode must be per_sqm or manual.",
    "STATUS_INVALID": "The status is not an allowed value.",
    "BOOLEAN_INVALID": "The value must be yes/no (true/false, 1/0).",
    "NUMBER_INVALID": "The value is not a finite number with at most 2 decimals.",
    "NUMBER_NEGATIVE": "The value can't be negative.",
    "AMOUNT_NOT_POSITIVE": "The amount must be greater than zero.",
    "DATE_INVALID": "The date is not a valid date.",
    "MONTH_INVALID": "The month must be YYYY-MM (or an Excel date).",
    "REQUIRED_VALUE_MISSING": "A value required for a new record is blank.",
    "TEXT_TOO_LONG": "The text is longer than the field allows.",
    "READING_DECREASES": "The current reading is lower than the previous reading.",
    # people
    "PERSON_RELATIONSHIP_UNCONFIRMED": "It is not confirmed that each cell names exactly one person.",
    "PERSON_MULTIPLE_NAMES": "The cell may name more than one person (/, &, comma, 'and').",
    "PERSON_CLEAR_NOT_SUPPORTED": "Ending an owner/tenancy through an import is not supported.",
    "OWNER_EQUALS_TENANT": "Owner and tenant cells are the same; the rule for this case is not approved.",
    "TENANT_WITHOUT_OWNER": "A tenant is given but no owner.",
    "POSSIBLE_DUPLICATE_PERSON": "The unit already has a person with the same name; names are never merged automatically.",
    "PERSON_NAME_REPEATED": "The same name appears on several properties; they are not merged (information only).",
    # relationships and collisions
    "UNRESOLVED_UNIT_REFERENCE": "The unit reference doesn't match any unit in the workbook or the crosswalk.",
    "UNRESOLVED_BILL_REFERENCE": "The billing reference doesn't match any bill in the workbook or the crosswalk.",
    "UNRESOLVED_ASSET_REFERENCE": "The parking/storage reference doesn't match a PARKING/STORAGE unit.",
    "UNIT_MONTH_DUPLICATE": "The same unit and month appear twice in this sheet.",
    "UNIT_MONTH_EXISTS": "CITYLAND9 already has a record for this unit and month that is not mapped to this row.",
    "EMPLOYEE_NO_MISSING": "The employee number is blank.",
    "EMPLOYEE_NO_DUPLICATE": "The same employee number appears twice in this sheet.",
    "EMPLOYEE_NO_CONFLICT": "The crosswalk and the employee number point to different employees.",
    "NO_NATURAL_KEY": "This record type has no natural key; without a crosswalk entry a re-run would add it again.",
    "PARENT_RECORD_NOT_VALID": "The unit or bill this row belongs to has its own problem; fix that row first.",
    "NEW_RECORD_NOT_ALLOWED": "The id is not a record of this database; adding records through this import is locked.",
    "ID_POINTS_TO_OTHER_RECORD": "The id belongs to a record of another unit or month here; the import would move it.",
    "BLANK_WOULD_RESET": "The cell is blank, and the import would replace the saved value with a default.",
    "LAYOUT_NOT_IMPORTABLE": "Only CITYLAND9's own export of this database can be imported here; other layouts are check-only.",
    # financial protections
    "PERIOD_CLOSED": "The month is in a closed period (books closed); it can't be changed by an import.",
    "BILL_ISSUED": "The bill was already issued; changing an issued bill needs a separately approved correction.",
    "BILL_CORRECTED_BY_HAND": "The bill was corrected by hand; an import can't change it.",
    "PAYMENT_TOTAL_MISMATCH": "The bill's amount paid differs from the sum of its payment rows.",
    "PAYMENT_DETAIL_MISSING": "The bill shows an amount paid but has no payment rows; receipts are never invented.",
    "STATUS_INCONSISTENT": "The status doesn't match the amounts (paid vs balance).",
    "OVERPAYMENT_WITHOUT_ADVANCE_LEDGER": "More was paid than charged, but the source has no advance/credit ledger.",
    "OPENING_ARREARS_NOT_CARRIED": "The first imported bill of the unit has a previous balance; CITYLAND9 recalculates arrears from earlier bills, so this amount would be lost.",
    "PREVIOUS_BALANCE_WITH_HISTORY": "The bill has a previous balance and earlier bills are also imported; the carried balance must reconcile, not be added twice.",
    "PENALTY_RECALCULATED": "CITYLAND9 recalculates penalties of normal bills; the imported penalty would be replaced.",
    "BILL_AFTER_CUTOFF": "The month is after the approved cutoff.",
    "BILL_OUTSIDE_SCOPE": "The month is outside the approved migration scope.",
    "HISTORY_IN_OPENING_BALANCE_MODE": "Opening-balance mode is configured, but bills before the cutoff month are included.",
    # readiness (not per row)
    "FINANCIAL_SOURCE_MISSING": "The workbook has no bills or payments, but historical bills are to be migrated.",
    "FINANCIAL_NOT_APPROVED": "The financial migration settings are not approved.",
    "FINANCIAL_CUTOFF_MISSING": "No approved cutoff date.",
    "FINANCIAL_MODE_MISSING": "The migration mode (full history or opening balances) is not decided.",
    "FINANCIAL_SCOPE_MISSING": "The scope (months and record types) is not defined.",
    "CONTROL_TOTALS_MISSING": "Source control totals are not provided.",
    "CONTROL_TOTALS_MISMATCH": "The workbook doesn't match the source control totals.",
    "ADVANCE_LEDGER_MISSING": "The source has no advance-payment/credit ledger.",
    "RECEIPTS_NOT_IN_SOURCE": "The source has no official receipts; they are never invented by an import.",
    "OPENING_BALANCE_UNSUPPORTED": "CITYLAND9 has no approved way to record an opening balance yet.",
    "IMPORTER_NOT_ENABLED": "Importing client records is not enabled in this version (check only).",
}


class CheckError(Exception):
    """The workbook can't be checked at all (not a workbook, too large)."""

    def __init__(self, code):
        super().__init__(MESSAGES[code])
        self.code = code


# ------------------------------------------------------------------ no-write guard
_GUARD = threading.local()
_READ_HEADS = ("SELECT", "SHOW", "PRAGMA", "WITH", "ROLLBACK", "BEGIN", "SAVEPOINT", "RELEASE", "SET", "DESCRIBE", "EXPLAIN")
_installed = set()


class DryRunWriteError(RuntimeError):
    pass


def _before_cursor_execute(conn, cursor, statement, parameters, context, executemany):
    if getattr(_GUARD, "active", False):
        head = statement.lstrip().split(None, 1)[0].upper() if statement.strip() else ""
        if head not in _READ_HEADS:
            raise DryRunWriteError("A database write was attempted during a dry-run check; it was refused.")


def _refuse_flush(session, flush_context, instances):
    if getattr(_GUARD, "active", False) and (session.new or session.dirty or session.deleted):
        raise DryRunWriteError("A database write was attempted during a dry-run check; it was refused.")


@contextmanager
def no_writes(db):
    """Refuse every write statement on this thread while the block runs; roll back afterwards."""
    from sqlalchemy import event
    engine = db.engine
    if id(engine) not in _installed:
        event.listen(engine, "before_cursor_execute", _before_cursor_execute)
        _installed.add(id(engine))
    if "session" not in _installed:
        event.listen(db.session, "before_flush", _refuse_flush)
        _installed.add("session")
    _GUARD.active = True
    try:
        yield
    finally:
        _GUARD.active = False
        db.session.rollback()


# ------------------------------------------------------------------ workbook reading
def read_workbook(stream, *, max_unpacked_mb, max_rows, max_columns):
    """{sheet: (headers, [(excel_row, values tuple)])}, same limits as the import."""
    from openpyxl import load_workbook
    stream.seek(0)
    try:
        with zipfile.ZipFile(stream) as archive:
            parts = archive.infolist()
            unpacked = sum(p.file_size for p in parts)
    except zipfile.BadZipFile:
        raise CheckError("FILE_NOT_XLSX") from None
    if len(parts) > 2000 or unpacked > max_unpacked_mb * 1024 * 1024:
        raise CheckError("FILE_TOO_LARGE")
    stream.seek(0)
    try:
        wb = load_workbook(stream, read_only=True, data_only=True)
    except Exception:
        raise CheckError("FILE_NOT_XLSX") from None
    sheets, limits = {}, []
    try:
        for ws in wb.worksheets:
            ws.reset_dimensions()   # a wrong stored dimension would hide columns
            iterator = ws.iter_rows(values_only=True)
            try:
                header = next(iterator)
            except StopIteration:
                sheets[ws.title] = ([], [])
                continue
            headers = [str(h).strip() if h is not None else "" for h in header]
            if len(headers) > max_columns:
                limits.append((ws.title, "SHEET_TOO_MANY_COLUMNS"))
                continue
            rows = []
            for n, values in enumerate(iterator, start=2):
                if any(v not in (None, "") for v in values):
                    if len(rows) >= max_rows:
                        limits.append((ws.title, "SHEET_TOO_MANY_ROWS"))
                        break
                    rows.append((n, tuple(values) + (None,) * max(0, len(headers) - len(values))))
            sheets[ws.title] = (headers, rows)
    finally:
        wb.close()
    return sheets, limits


# ------------------------------------------------------------------ cell helpers
def state_of(value, present=True):
    """missing_column | blank | clear | zero | value."""
    if not present:
        return "missing_column"
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return "blank"
    if isinstance(value, str) and value.strip().upper() == CLEAR_MARKER:
        return "clear"
    if isinstance(value, bool):
        return "value"
    if isinstance(value, (int, float)) and value == 0:
        return "zero"
    if isinstance(value, str) and re.fullmatch(r"0+(\.0+)?", value.strip()):
        return "zero"
    return "value"


def norm_key(text):
    return re.sub(r"\s+", "", str(text)).upper()


def norm_name(text):
    return re.sub(r"\s+", " ", str(text)).strip().upper()


def external_id_text(value):
    """Canonical text of an external id, or None when it isn't an id."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value == value and value not in (float("inf"), float("-inf")) and value.is_integer() else None
    if isinstance(value, str) and value.strip():
        return value.strip()[:100] if len(value.strip()) <= 100 else None
    return None


def parse_number(value, *, places=2):
    """Decimal, or raise InputError. Rejects NaN/inf, text, more than `places` decimals."""
    if isinstance(value, bool):
        raise InputError("value", "not a number")
    if isinstance(value, float) and (value != value or value in (float("inf"), float("-inf"))):
        raise InputError("value", "not finite")
    text = repr(value) if isinstance(value, float) else str(value)
    return parse_decimal(text, label="value", places=places)


def parse_month(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m")
    if isinstance(value, date):
        return value.strftime("%Y-%m")
    text = str(value).strip()
    return text if MONTH_RE.match(text) else None


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def parse_bool(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    text = str(value).strip().lower()
    if text in ("true", "yes", "y", "1"):
        return True
    if text in ("false", "no", "n", "0"):
        return False
    return None


# ------------------------------------------------------------------ report
class Report:
    def __init__(self):
        self.issues = []
        self.row_rank = {}            # (sheet, row) -> worst rank
        self.rows_seen = {}           # sheet -> set(rows)
        self.proposals = {}           # entity -> {insert, update, link, unchanged}
        self.cell_states = {}         # sheet -> column -> {state: n}
        self.readiness = {}           # area -> list of codes (blockers)
        self.notes = []

    def see(self, sheet, row):
        self.rows_seen.setdefault(sheet, set()).add(row)

    def add(self, sheet, row, column, code, severity):
        self.issues.append({"sheet": sheet, "row": row, "column": column, "code": code, "severity": severity,
                            "message": MESSAGES[code]})
        if row is not None:
            key = (sheet, row)
            self.row_rank[key] = max(self.row_rank.get(key, -1), RANK[severity])

    def count_state(self, sheet, column, state):
        cols = self.cell_states.setdefault(sheet, {})
        bucket = cols.setdefault(column, {"missing_column": 0, "blank": 0, "clear": 0, "zero": 0, "value": 0})
        bucket[state] += 1

    def propose(self, entity, action):
        self.proposals.setdefault(entity, {"insert": 0, "update": 0, "link": 0, "unchanged": 0})[action] += 1

    def block_area(self, area, code):
        if code not in self.readiness.setdefault(area, []):
            self.readiness[area].append(code)

    def row_status(self, sheet, row):
        rank = self.row_rank.get((sheet, row), -1)
        return {3: "blocked", 2: "conflict", 1: "unresolved"}.get(rank, "valid")

    def as_dict(self, *, fmt, file_info, config_problems):
        sheets = {}
        for sheet, rows in self.rows_seen.items():
            counts = {"rows": len(rows), "valid": 0, "blocked": 0, "conflict": 0, "unresolved": 0}
            for row in rows:
                counts[self.row_status(sheet, row)] += 1
            sheets[sheet] = counts
        codes = {}
        for i in self.issues:
            codes.setdefault(i["code"], {"severity": i["severity"], "count": 0, "message": i["message"]})["count"] += 1
        readiness = {area: {"ready": not codes_, "blockers": [{"code": c, "message": MESSAGES[c]} for c in codes_]}
                     for area, codes_ in self.readiness.items()}
        ordered = sorted(self.issues, key=lambda i: (-RANK[i["severity"]], i["sheet"] or "", i["row"] or 0, i["column"] or ""))
        return {
            "format": fmt,
            "dryRun": True,
            "databaseWrites": 0,
            "file": file_info,
            "configProblems": [{"code": c, "message": MESSAGES[c]} for c in config_problems],
            "sheets": sheets,
            "proposals": self.proposals,
            "cellStates": self.cell_states,
            "issueCounts": codes,
            "issues": ordered[:MAX_ISSUES_LISTED],
            "issuesTruncated": len(ordered) > MAX_ISSUES_LISTED,
            "readiness": readiness,
            "readyForImport": False,   # stage: check only; importing client records is not enabled
            "notes": self.notes,
        }


# ------------------------------------------------------------------ database view (read only)
class DbView:
    """What the checks need from CITYLAND9, loaded with SELECTs only."""

    def __init__(self, legacy, source_code, roundtrip=False):
        m = legacy
        self.roundtrip = roundtrip
        self.units = {u.id: u for u in m.Unit.query.all()}
        self.units_by_key = {}
        for u in self.units.values():
            self.units_by_key.setdefault(norm_key(u.unit_no), []).append(u)
        self.bills = {(b.unit_id, b.billing_month): b for b in m.Billing.query.all()}
        self.readings = {(w.unit_id, w.reading_month): w for w in m.WaterReading.query.all()}
        self.employees_by_no = {norm_key(e.employee_no): e for e in m.Employee.query.all()}
        self.people_by_unit = {}
        for cls, kind in ((m.Owner, "owner"), (m.Tenant, "tenant")):
            name_field = "owner_name" if kind == "owner" else "tenant_name"
            for p in cls.query.all():
                self.people_by_unit.setdefault((kind, p.unit_id), set()).add(norm_name(getattr(p, name_field) or ""))
        self.closed_through = m.books_closed_through()
        self.source = None
        self.crosswalk = {}       # (entity, external) -> target
        self.crosswalk_rev = {}   # (entity, target) -> external
        self.crosswalk_ok = True
        self.by_id = {"unit": self.units}
        self.water_rate = Decimal(str(m.setting_float("water_rate", 50)))
        if roundtrip:
            # CITYLAND9's own export: its ids ARE this database's ids (identity mapping, nothing stored).
            for entity, model in (("owner", m.Owner), ("tenant", m.Tenant), ("billing", m.Billing), ("payment", m.Payment),
                                  ("water_reading", m.WaterReading), ("employee", m.Employee), ("expense", m.Expense)):
                self.by_id[entity] = {r.id: r for r in model.query.all()}
            for entity, records in self.by_id.items():
                for rid in records:
                    self.crosswalk[(entity, str(rid))] = rid
                    self.crosswalk_rev[(entity, rid)] = str(rid)
            return
        try:
            if source_code:
                self.source = m.ImportSource.query.filter_by(code=source_code).first()
            if self.source:
                for x in m.ImportCrosswalk.query.filter_by(source_id=self.source.id).all():
                    self.crosswalk[(x.entity_type, x.external_id)] = x.target_id
                    self.crosswalk_rev[(x.entity_type, x.target_id)] = x.external_id
                for entity, model in (("owner", m.Owner), ("tenant", m.Tenant), ("billing", m.Billing), ("payment", m.Payment),
                                      ("water_reading", m.WaterReading), ("employee", m.Employee), ("expense", m.Expense)):
                    ids = {t for (e, _), t in self.crosswalk.items() if e == entity}
                    if ids:
                        self.by_id[entity] = {r.id: r for r in model.query.filter(model.id.in_(ids)).all()}
        except Exception:
            self.crosswalk_ok = False
            m.db.session.rollback()


# ------------------------------------------------------------------ configuration
def load_config(path):
    if not path or not os.path.exists(path):
        return {}
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def source_code_of(config):
    code = str(((config or {}).get("source_system") or {}).get("code") or "").strip()
    return code if code and not code.startswith("<") else ""


def financial_blockers(config, report):
    fin = (config or {}).get("financial") or {}
    if fin.get("approved") is not True or not fin.get("approved_by") or not fin.get("approved_on"):
        report.block_area("financial", "FINANCIAL_NOT_APPROVED")
    cutoff = parse_date(fin.get("cutoff_date")) if fin.get("cutoff_date") else None
    if not cutoff:
        report.block_area("financial", "FINANCIAL_CUTOFF_MISSING")
    if fin.get("mode") not in FINANCIAL_MODES:
        report.block_area("financial", "FINANCIAL_MODE_MISSING")
    scope = fin.get("scope") or {}
    if not (scope.get("entities") and parse_month(scope.get("first_month") or "") and parse_month(scope.get("last_month") or "")):
        report.block_area("financial", "FINANCIAL_SCOPE_MISSING")
    totals = fin.get("control_totals") or {}
    if any(totals.get(k) in (None, "") for k in ("billing_rows", "billing_charges_total", "payment_rows", "payment_total")):
        report.block_area("financial", "CONTROL_TOTALS_MISSING")
    return fin, cutoff


# ------------------------------------------------------------------ format: properties
def check_properties(sheets, config, db, report):
    cfg = (config or {}).get("properties") or {}
    sheet = cfg.get("sheet") or "Properties"
    cols = {"external_id": "ID", "unit_no": "Name", "area_sqm": "Size", "group": "Group", "owner": "Owner", "tenant": "Tenant"}
    cols.update(cfg.get("columns") or {})
    headers, rows = sheets[sheet]
    index = {h: i for i, h in enumerate(headers) if h}
    pos = {k: index.get(v) for k, v in cols.items()}
    for key in ("external_id", "unit_no", "area_sqm", "group"):
        if pos[key] is None:
            report.add(sheet, None, cols[key], "REQUIRED_COLUMN_MISSING", "blocked")
            report.block_area("properties", "REQUIRED_COLUMN_MISSING")
    if pos["external_id"] is None or pos["unit_no"] is None:
        for n, _ in rows:
            report.see(sheet, n)
            report.add(sheet, n, None, "REQUIRED_COLUMN_MISSING", "blocked")
        return
    source_ok = bool(db.source) or bool(source_code_of(config))
    unit_no_ok = cfg.get("unit_no_confirmed") is True
    groups = cfg.get("group_mapping") or {}
    people_ok = ((cfg.get("people") or {}).get("one_person_per_cell_confirmed") is True)

    def v(values, key):
        return values[pos[key]] if pos[key] is not None else None

    ext_count, key_count, name_rows = {}, {}, {}
    for n, values in rows:
        ext = external_id_text(v(values, "external_id"))
        if ext:
            ext_count[ext] = ext_count.get(ext, 0) + 1
        unit_no = str(v(values, "unit_no") or "").strip()
        if unit_no:
            key_count[norm_key(unit_no)] = key_count.get(norm_key(unit_no), 0) + 1
        for key in ("owner", "tenant"):
            text = v(values, key)
            if state_of(text, pos[key] is not None) == "value":
                name_rows.setdefault(norm_name(text), set()).add(n)

    for n, values in rows:
        report.see(sheet, n)
        for key in cols:
            report.count_state(sheet, cols[key], state_of(v(values, key), pos[key] is not None))
        # --- identity
        raw_ext = v(values, "external_id")
        ext = external_id_text(raw_ext)
        unit_ok = True
        if state_of(raw_ext) in ("blank", "clear"):
            report.add(sheet, n, cols["external_id"], "EXTERNAL_ID_MISSING", "blocked"); unit_ok = False
        elif ext is None:
            report.add(sheet, n, cols["external_id"], "EXTERNAL_ID_FORMAT", "blocked"); unit_ok = False
        elif ext_count.get(ext, 0) > 1:
            report.add(sheet, n, cols["external_id"], "EXTERNAL_ID_DUPLICATE", "blocked"); unit_ok = False
        raw_no = v(values, "unit_no")
        unit_no = str(raw_no).strip() if state_of(raw_no) == "value" else ""
        if state_of(raw_no) == "clear":
            report.add(sheet, n, cols["unit_no"], "VALUE_CLEAR_NOT_ALLOWED", "blocked"); unit_ok = False
        elif not unit_no:
            report.add(sheet, n, cols["unit_no"], "UNIT_NO_MISSING", "blocked"); unit_ok = False
        elif len(unit_no) > 40:
            report.add(sheet, n, cols["unit_no"], "UNIT_NO_TOO_LONG", "blocked"); unit_ok = False
        elif key_count.get(norm_key(unit_no), 0) > 1:
            report.add(sheet, n, cols["unit_no"], "UNIT_NO_DUPLICATE", "blocked"); unit_ok = False
        if not unit_no_ok:
            report.add(sheet, n, cols["unit_no"], "UNIT_NO_UNCONFIRMED", "unresolved"); unit_ok = False
        if not source_ok:
            unit_ok = False

        # --- match (only with a confirmed unit-number rule and a source namespace)
        target, action = None, None
        if unit_ok:
            mapped = db.crosswalk.get(("unit", ext))
            candidates = db.units_by_key.get(norm_key(unit_no), [])
            if mapped is not None:
                target = db.units.get(mapped)
                if target is None:
                    report.add(sheet, n, cols["external_id"], "CROSSWALK_TARGET_MISSING", "conflict"); unit_ok = False
                elif norm_key(target.unit_no) != norm_key(unit_no):
                    report.add(sheet, n, cols["unit_no"], "CROSSWALK_UNIT_NO_MISMATCH", "conflict"); unit_ok = False
                else:
                    action = "update"
            elif len(candidates) > 1:
                report.add(sheet, n, cols["unit_no"], "AMBIGUOUS_UNIT_MATCH", "conflict"); unit_ok = False
            elif candidates:
                target = candidates[0]
                other = db.crosswalk_rev.get(("unit", target.id))
                if other is not None and other != ext:
                    report.add(sheet, n, cols["unit_no"], "UNIT_MAPPED_TO_OTHER_EXTERNAL_ID", "conflict"); unit_ok = False
                else:
                    action = "link"
            else:
                action = "insert"

        # --- size (confirmed sqm)
        raw_area = v(values, "area_sqm")
        area_state = state_of(raw_area, pos["area_sqm"] is not None)
        area = None
        if area_state == "clear":
            report.add(sheet, n, cols["area_sqm"], "VALUE_CLEAR_NOT_ALLOWED", "blocked"); unit_ok = False
        elif area_state in ("missing_column", "blank"):
            if action == "insert" or target is None:
                report.add(sheet, n, cols["area_sqm"], "AREA_REQUIRED", "blocked"); unit_ok = False
        elif area_state == "zero":
            report.add(sheet, n, cols["area_sqm"], "AREA_ZERO", "blocked"); unit_ok = False
        else:
            try:
                area = parse_number(raw_area, places=2)
                if area < 0:
                    report.add(sheet, n, cols["area_sqm"], "AREA_NEGATIVE", "blocked"); unit_ok = False
            except InputError as exc:
                code = "AREA_PRECISION" if "decimal" in exc.message else "AREA_NOT_A_NUMBER"
                report.add(sheet, n, cols["area_sqm"], code, "blocked"); unit_ok = False

        # --- group
        raw_group = v(values, "group")
        group_state = state_of(raw_group, pos["group"] is not None)
        mapping = None
        if group_state in ("blank", "clear", "missing_column"):
            report.add(sheet, n, cols["group"], "GROUP_MISSING", "blocked"); unit_ok = False
        else:
            label = str(raw_group).strip()
            mapping = groups.get(label) or groups.get(label.upper())
            if mapping is None:
                report.add(sheet, n, cols["group"], "GROUP_UNMAPPED", "unresolved"); unit_ok = False
            elif mapping.get("approved") is not True:
                report.add(sheet, n, cols["group"], "GROUP_UNCONFIRMED", "unresolved"); unit_ok = False
            elif mapping.get("unit_type") and mapping["unit_type"] not in UNIT_TYPES:
                report.add(sheet, n, cols["group"], "UNIT_TYPE_INVALID", "blocked"); unit_ok = False

        if unit_ok and action:
            if action == "update" or action == "link":
                same_area = area is None or Decimal(str(target.area_sqm or 0)).quantize(Decimal("0.01")) == area
                same_floor = not (mapping or {}).get("floor") or (target.floor or "") == mapping["floor"]
                same_type = not (mapping or {}).get("unit_type") or (target.unit_type or "") == mapping["unit_type"]
                report.propose("unit", action if action == "link" else ("unchanged" if same_area and same_floor and same_type else "update"))
            else:
                report.propose("unit", "insert")

        # --- people (never merged by name, never accounts)
        owner_raw, tenant_raw = v(values, "owner"), v(values, "tenant")
        owner_state = state_of(owner_raw, pos["owner"] is not None)
        tenant_state = state_of(tenant_raw, pos["tenant"] is not None)
        for key, raw, st in (("owner", owner_raw, owner_state), ("tenant", tenant_raw, tenant_state)):
            if pos[key] is None:
                continue
            if st == "clear":
                report.add(sheet, n, cols[key], "PERSON_CLEAR_NOT_SUPPORTED", "unresolved")
            elif st in ("value", "zero"):
                text = str(raw)
                if MULTI_PERSON.search(text):
                    report.add(sheet, n, cols[key], "PERSON_MULTIPLE_NAMES", "unresolved")
                elif not people_ok:
                    report.add(sheet, n, cols[key], "PERSON_RELATIONSHIP_UNCONFIRMED", "unresolved")
                if len(text.strip()) > 200:
                    report.add(sheet, n, cols[key], "TEXT_TOO_LONG", "blocked")
                if len(name_rows.get(norm_name(text), ())) > 1:
                    report.add(sheet, n, cols[key], "PERSON_NAME_REPEATED", "warning")
                if target is not None and norm_name(text) in db.people_by_unit.get((key, target.id), set()):
                    report.add(sheet, n, cols[key], "POSSIBLE_DUPLICATE_PERSON", "unresolved")
        if owner_state == "value" and tenant_state == "value" and norm_name(owner_raw) == norm_name(tenant_raw):
            rule = ((cfg.get("people") or {}).get("owner_equals_tenant_rule"))
            if not rule:
                report.add(sheet, n, cols["tenant"], "OWNER_EQUALS_TENANT", "unresolved")
        if tenant_state == "value" and owner_state in ("blank", "missing_column"):
            report.add(sheet, n, cols["tenant"], "TENANT_WITHOUT_OWNER", "unresolved")

    # readiness
    statuses = [report.row_status(sheet, n) for n, _ in rows]
    if not source_ok:
        report.block_area("properties", "SOURCE_SYSTEM_NOT_CONFIGURED")
    if not unit_no_ok:
        report.block_area("properties", "UNIT_NO_UNCONFIRMED")
    people_columns = {cols["owner"], cols["tenant"]}
    report.readiness.setdefault("properties", [])
    report.readiness.setdefault("people", [])
    for i in report.issues:
        if i["sheet"] != sheet or i["severity"] == "warning":
            continue
        report.block_area("people" if i["column"] in people_columns else "properties", i["code"])
    # This layout carries no bills or payments: the historical-bill migration needs another source.
    report.block_area("financial", "FINANCIAL_SOURCE_MISSING")
    financial_blockers(config, report)
    report.notes.append(f"{statuses.count('valid')} of {len(rows)} rows are valid; people are never proposed until "
                        "their relationship rules are approved.")


# ------------------------------------------------------------------ format: CITYLAND9 export
EXPORT_SHEETS = ("Units", "Owners", "Tenants", "Billing", "Payments", "WaterReadings", "Employees", "Expenses")
BILL_CHARGES = ("assessment", "parking_dues", "storage_dues", "water", "other", "penalty", "adjustment")
CENT = Decimal("0.01")


def _dec(value):
    return Decimal(str(value or 0)).quantize(CENT)


def check_export(sheets, config, db, report):
    """CITYLAND9's export layout. Two modes:
    migration           ids are another system's ids (crosswalk namespace of the configured source)
    cityland_roundtrip  ids are THIS database's own ids (the file is CITYLAND9's own export, edited):
                        every row must update an existing record whose id matches; nothing new is
                        added; a row whose values are unchanged is not blocked by closed periods or
                        issued bills (the import would not change it)."""
    rt = db.roundtrip
    source_ok = rt or bool(db.source) or bool(source_code_of(config))
    if rt:
        fin, cutoff = {}, None
    else:
        fin, cutoff = financial_blockers(config, report)
    cutoff_month = cutoff.strftime("%Y-%m") if cutoff else None
    scope = fin.get("scope") or {}
    first_m, last_m = parse_month(scope.get("first_month") or ""), parse_month(scope.get("last_month") or "")
    mode = fin.get("mode")
    if not source_ok:
        report.block_area("records", "SOURCE_SYSTEM_NOT_CONFIGURED")
    if "Units" not in sheets and not rt:
        report.add("Units", None, None, "REQUIRED_SHEET_MISSING", "blocked")
        report.block_area("records", "REQUIRED_SHEET_MISSING")

    def table(name):
        headers, rows = sheets.get(name, ([], []))
        idx = {h: i for i, h in enumerate(headers) if h}
        return idx, rows

    def get(idx, values, col):
        return values[idx[col]] if col in idx else None

    def has(idx, col):
        return col in idx

    def ext_of(idx, values):
        return external_id_text(get(idx, values, "id")) if has(idx, "id") else None

    def check_ext_column(idx, rows):
        seen = {}
        for n, values in rows:
            ext = ext_of(idx, values)
            if ext:
                seen.setdefault(ext, []).append(n)
        return {n for ns in seen.values() if len(ns) > 1 for n in ns}

    def blank(idx, values, col):
        return state_of(get(idx, values, col), has(idx, col)) in ("missing_column", "blank")

    def number(sheet, n, col, idx, values, *, negative=False, positive=False, required=False):
        raw = get(idx, values, col)
        st = state_of(raw, has(idx, col))
        report.count_state(sheet, col, st)
        if st in ("missing_column", "blank"):
            if required:
                report.add(sheet, n, col, "REQUIRED_VALUE_MISSING", "blocked")
            return None
        if st == "clear":
            if required:
                report.add(sheet, n, col, "VALUE_CLEAR_NOT_ALLOWED", "blocked")
            return None
        try:
            value = parse_number(raw, places=2)
        except InputError:
            report.add(sheet, n, col, "NUMBER_INVALID", "blocked")
            return None
        if value < 0 and not negative:
            report.add(sheet, n, col, "NUMBER_NEGATIVE", "blocked")
        if positive and value <= 0:
            report.add(sheet, n, col, "AMOUNT_NOT_POSITIVE", "blocked")
        return value

    def choice(sheet, n, col, idx, values, allowed, code, required=False, current=None):
        raw = get(idx, values, col)
        st = state_of(raw, has(idx, col))
        report.count_state(sheet, col, st)
        if st in ("missing_column", "blank"):
            if required:
                report.add(sheet, n, col, "REQUIRED_VALUE_MISSING", "blocked")
            return None
        if st == "clear":
            report.add(sheet, n, col, "VALUE_CLEAR_NOT_ALLOWED", "blocked")
            return None
        text = str(raw).strip()
        if text not in allowed and not (current is not None and text == str(current).strip()):
            report.add(sheet, n, col, code, "blocked")   # an old value kept unchanged is not re-judged
            return None
        return text

    def a_date(sheet, n, col, idx, values, required=False):
        raw = get(idx, values, col)
        st = state_of(raw, has(idx, col))
        report.count_state(sheet, col, st)
        if st in ("missing_column", "blank", "clear"):
            if required:
                report.add(sheet, n, col, "REQUIRED_VALUE_MISSING", "blocked")
            return None
        d = parse_date(raw)
        if d is None:
            report.add(sheet, n, col, "DATE_INVALID", "blocked")
        return d

    def a_bool(sheet, n, col, idx, values):
        raw = get(idx, values, col)
        st = state_of(raw, has(idx, col))
        report.count_state(sheet, col, st)
        if st in ("missing_column", "blank"):
            return None
        value = parse_bool(raw)
        if value is None:
            report.add(sheet, n, col, "BOOLEAN_INVALID", "blocked")
        return value

    def a_month(sheet, n, col, idx, values):
        raw = get(idx, values, col)
        st = state_of(raw, has(idx, col))
        report.count_state(sheet, col, st)
        if st in ("missing_column", "blank", "clear"):
            report.add(sheet, n, col, "REQUIRED_VALUE_MISSING", "blocked")
            return None
        month = parse_month(raw)
        if month is None:
            report.add(sheet, n, col, "MONTH_INVALID", "blocked")
        return month

    def text_field(sheet, n, col, idx, values, limit, required=False):
        raw = get(idx, values, col)
        st = state_of(raw, has(idx, col))
        report.count_state(sheet, col, st)
        if st in ("missing_column", "blank", "clear"):
            if required:
                report.add(sheet, n, col, "VALUE_CLEAR_NOT_ALLOWED" if st == "clear" else "REQUIRED_VALUE_MISSING", "blocked")
            return None
        text = str(raw).strip()
        if len(text) > limit:
            report.add(sheet, n, col, "TEXT_TOO_LONG", "blocked")
        return text

    def closed(month):
        return bool(db.closed_through and month and month <= db.closed_through)

    def existing(sheet, n, entity, ext):
        """Round trip: the record this row updates (its id must exist here); migration: crosswalk hit or None."""
        target_id = db.crosswalk.get((entity, ext)) if ext else None
        record = db.by_id.get(entity, {}).get(target_id) if target_id is not None else None
        if rt and ext and record is None:
            report.add(sheet, n, "id", "NEW_RECORD_NOT_ALLOWED", "blocked")
        return record

    def reset_check(sheet, n, idx, values, record, col, current, reset_value):
        """The classic import replaces a blank/missing cell with a default: block if that changes the record."""
        if record is not None and blank(idx, values, col) and current != reset_value:
            report.add(sheet, n, col, "BLANK_WOULD_RESET", "blocked")
            return True
        return False

    def differs(present, current):
        return present is not None and present != current

    def external_text(value):
        """A text cell as the classic import would store it (blank keeps the saved value: None here)."""
        return str(value).strip() if state_of(value) == "value" else None

    # ---------- Units
    idx, rows = table("Units")
    dup = check_ext_column(idx, rows)
    key_count = {}
    for _, values in rows:
        no = str(get(idx, values, "unit_no") or "").strip()
        if no:
            key_count[norm_key(no)] = key_count.get(norm_key(no), 0) + 1
    unit_ref = {}       # workbook unit id -> ("db", unit) | ("new", None) | None (invalid)
    unit_types = {}
    for n, values in rows:
        report.see("Units", n)
        ok = True
        ext = ext_of(idx, values)
        if not has(idx, "id") or ext is None:
            report.add("Units", n, "id", "EXTERNAL_ID_MISSING", "blocked"); ok = False
        elif n in dup:
            report.add("Units", n, "id", "EXTERNAL_ID_DUPLICATE", "blocked"); ok = False
        unit_no = text_field("Units", n, "unit_no", idx, values, 40, required=True)
        if unit_no and key_count.get(norm_key(unit_no), 0) > 1:
            report.add("Units", n, "unit_no", "UNIT_NO_DUPLICATE", "blocked"); ok = False
        area = number("Units", n, "area_sqm", idx, values)
        target, action = None, None
        if unit_no and ext and source_ok:
            mapped = db.crosswalk.get(("unit", ext))
            candidates = db.units_by_key.get(norm_key(unit_no), [])
            if mapped is not None:
                target = db.units.get(mapped)
                if target is None:
                    report.add("Units", n, "id", "CROSSWALK_TARGET_MISSING", "conflict"); ok = False
                elif norm_key(target.unit_no) != norm_key(unit_no):
                    report.add("Units", n, "unit_no", "CROSSWALK_UNIT_NO_MISMATCH", "conflict"); ok = False
                else:
                    action = "update"
            elif len(candidates) > 1:
                report.add("Units", n, "unit_no", "AMBIGUOUS_UNIT_MATCH", "conflict"); ok = False
            elif candidates:
                target = candidates[0]
                other = db.crosswalk_rev.get(("unit", target.id))
                if other is not None and other != ext:
                    report.add("Units", n, "unit_no", "UNIT_MAPPED_TO_OTHER_EXTERNAL_ID", "conflict"); ok = False
                else:
                    action = "link"
            elif rt:
                report.add("Units", n, "id", "NEW_RECORD_NOT_ALLOWED", "blocked"); ok = False
            else:
                action = "insert"
                if area is None or area <= 0:
                    report.add("Units", n, "area_sqm", "AREA_REQUIRED" if area is None else "AREA_ZERO", "blocked"); ok = False
        cur = target if action in ("update",) else None
        utype = choice("Units", n, "unit_type", idx, values, UNIT_TYPES, "UNIT_TYPE_INVALID", current=getattr(cur, "unit_type", None))
        dues = choice("Units", n, "dues_mode", idx, values, DUES_MODES, "DUES_MODE_INVALID", current=getattr(cur, "dues_mode", None))
        status = choice("Units", n, "status", idx, values, UNIT_STATUSES, "STATUS_INVALID", current=getattr(cur, "status", None))
        floor = text_field("Units", n, "floor", idx, values, 20)
        occupancy = text_field("Units", n, "occupancy_type", idx, values, 20)
        rate = number("Units", n, "unit_rate_per_sqm", idx, values)
        mdues = number("Units", n, "manual_monthly_dues", idx, values)
        flags = {col: a_bool("Units", n, col, idx, values) for col in ("include_parking", "include_storage", "active")}
        if cur is not None:
            refs = {col: external_id_text(get(idx, values, col)) for col in ("assigned_parking_unit_id", "assigned_storage_unit_id")}
            unit_changed = any([
                reset_check("Units", n, idx, values, cur, "manual_monthly_dues", _dec(cur.manual_monthly_dues), _dec(0)),
                reset_check("Units", n, idx, values, cur, "active", bool(cur.active), True),
                # a blank assignment cell REMOVES the unit's parking/storage assignment in the classic import
                reset_check("Units", n, idx, values, cur, "assigned_parking_unit_id", cur.assigned_parking_unit_id, None),
                reset_check("Units", n, idx, values, cur, "assigned_storage_unit_id", cur.assigned_storage_unit_id, None),
                differs(utype, cur.unit_type), differs(dues, cur.dues_mode), differs(status, cur.status),
                differs(floor, cur.floor), differs(occupancy, cur.occupancy_type),
                differs(area, _dec(cur.area_sqm)), differs(rate, _dec(cur.unit_rate_per_sqm)),
                differs(mdues, _dec(cur.manual_monthly_dues))]
                + [differs(v, bool(getattr(cur, col))) for col, v in flags.items()]
                + [refs[col] is not None and refs[col] != str(getattr(cur, col) or "") for col in refs])
            if not unit_changed and action == "update":
                action = "unchanged"
            # The classic import turns an owner typed on the Units sheet into a NEW owner record when the
            # unit has none.
            if rt and state_of(get(idx, values, "owner_name"), has(idx, "owner_name")) == "value" \
                    and not db.people_by_unit.get(("owner", cur.id)):
                report.add("Units", n, "owner_name", "NEW_RECORD_NOT_ALLOWED", "blocked")
        if ext:
            ok = ok and report.row_status("Units", n) == "valid"
            unit_ref[ext] = (("db", target) if target is not None else ("new", None)) if ok else None
            unit_types[ext] = utype or (target.unit_type if target is not None else None)
        if ok and action and report.row_status("Units", n) == "valid":
            report.propose("unit", action)
    for n, values in rows:   # parking/storage assignments must resolve to a PARKING/STORAGE unit
        for col, kind in (("assigned_parking_unit_id", "PARKING"), ("assigned_storage_unit_id", "STORAGE")):
            raw = get(idx, values, col)
            if state_of(raw, has(idx, col)) == "value":
                ref = external_id_text(raw)
                ref_type = unit_types.get(ref)
                if ref not in unit_ref and rt and ref is not None and ref.isdigit() and int(ref) in db.units:
                    ref_type = db.units[int(ref)].unit_type     # round trip: an existing unit not in the sheet
                    if (ref_type or "").upper() == kind:
                        continue
                if ref not in unit_ref or unit_ref[ref] is None or (ref_type or "").upper() != kind:
                    report.add("Units", n, col, "UNRESOLVED_ASSET_REFERENCE", "blocked")

    def resolve_unit(sheet, n, idx, values):
        raw = get(idx, values, "unit_id")
        ref = external_id_text(raw)
        if ref is not None and ref in unit_ref and unit_ref[ref] is None:
            report.add(sheet, n, "unit_id", "PARENT_RECORD_NOT_VALID", "unresolved")
            return None
        if ref is not None and ref in unit_ref:
            return unit_ref[ref]
        if ref is not None and ("unit", ref) in db.crosswalk and db.units.get(db.crosswalk[("unit", ref)]) is not None:
            return ("db", db.units[db.crosswalk[("unit", ref)]])
        report.add(sheet, n, "unit_id", "UNRESOLVED_UNIT_REFERENCE", "blocked")
        return None

    def points_elsewhere(sheet, n, column, record, ref, extra=True):
        """The id names a record of ANOTHER unit (or month): the classic import would move it."""
        if record is not None and ref and ref[0] == "db" and (record.unit_id != ref[1].id or not extra):
            report.add(sheet, n, column, "ID_POINTS_TO_OTHER_RECORD", "conflict")

    # ---------- Owners / Tenants
    for sheet, entity, name_col, kind in (("Owners", "owner", "owner_name", "owner"), ("Tenants", "tenant", "tenant_name", "tenant")):
        idx, rows = table(sheet)
        dup = check_ext_column(idx, rows)
        for n, values in rows:
            report.see(sheet, n)
            ext = ext_of(idx, values)
            if ext is None:
                report.add(sheet, n, "id", "EXTERNAL_ID_MISSING", "blocked")
            elif n in dup:
                report.add(sheet, n, "id", "EXTERNAL_ID_DUPLICATE", "blocked")
            rec = existing(sheet, n, entity, ext)
            ref = resolve_unit(sheet, n, idx, values)
            points_elsewhere(sheet, n, "unit_id", rec, ref)
            name = text_field(sheet, n, name_col, idx, values, 200, required=True)
            text_field(sheet, n, "contact_no", idx, values, 80 if kind == "tenant" else 100)
            text_field(sheet, n, "email", idx, values, 160 if kind == "tenant" else 200)
            a_date(sheet, n, "move_in", idx, values)
            a_date(sheet, n, "move_out", idx, values)
            pstatus = choice(sheet, n, "status", idx, values, PERSON_STATUSES, "STATUS_INVALID", current=getattr(rec, "status", None))
            representative = a_bool(sheet, n, "representative", idx, values) if kind == "tenant" else None
            person_changed = rec is None or any([
                differs(name, getattr(rec, name_col)), differs(pstatus, rec.status),
                differs(external_text(get(idx, values, "contact_no")), rec.contact_no),
                differs(external_text(get(idx, values, "email")), rec.email),
                differs(external_text(get(idx, values, "notes")), rec.notes),
                differs(parse_date(get(idx, values, "move_in")) if state_of(get(idx, values, "move_in"), has(idx, "move_in")) == "value" else None, rec.move_in),
                differs(parse_date(get(idx, values, "move_out")) if state_of(get(idx, values, "move_out"), has(idx, "move_out")) == "value" else None, rec.move_out),
                differs(representative, bool(getattr(rec, "representative", False)) if kind == "tenant" else None)])
            unchanged_name = rec is not None and name is not None and norm_name(name) == norm_name(getattr(rec, name_col) or "")
            if name and MULTI_PERSON.search(name) and not unchanged_name:
                report.add(sheet, n, name_col, "PERSON_MULTIPLE_NAMES", "unresolved")
            mapped = db.crosswalk.get((entity, ext)) if ext else None
            if name and ref and ref[0] == "db" and mapped is None and norm_name(name) in db.people_by_unit.get((kind, ref[1].id), set()):
                report.add(sheet, n, name_col, "POSSIBLE_DUPLICATE_PERSON", "unresolved")
            if report.row_status(sheet, n) == "valid" and source_ok:
                report.propose(entity, ("update" if person_changed else "unchanged") if mapped is not None else "insert")

    # ---------- Water readings
    idx, rows = table("WaterReadings")
    seen, proposal = {}, {}
    dup = check_ext_column(idx, rows)
    for n, values in rows:
        report.see("WaterReadings", n)
        ext = ext_of(idx, values)
        if ext is None:
            report.add("WaterReadings", n, "id", "EXTERNAL_ID_MISSING", "blocked")
        elif n in dup:
            report.add("WaterReadings", n, "id", "EXTERNAL_ID_DUPLICATE", "blocked")
        rec = existing("WaterReadings", n, "water_reading", ext)
        ref = resolve_unit("WaterReadings", n, idx, values)
        month = a_month("WaterReadings", n, "reading_month", idx, values)
        points_elsewhere("WaterReadings", n, "reading_month", rec, ref, extra=(rec is None or rec.reading_month == month))
        prev = number("WaterReadings", n, "previous_reading", idx, values)
        cur = number("WaterReadings", n, "current_reading", idx, values)
        if prev is not None and cur is not None and cur < prev:
            report.add("WaterReadings", n, "current_reading", "READING_DECREASES", "blocked")
        rate = number("WaterReadings", n, "rate", idx, values)
        rdate = a_date("WaterReadings", n, "reading_date", idx, values)
        pdate = a_date("WaterReadings", n, "paid_date", idx, values)
        paid = a_bool("WaterReadings", n, "paid", idx, values)
        changed = rec is None
        if rec is not None:
            changed = any([
                reset_check("WaterReadings", n, idx, values, rec, "previous_reading", _dec(rec.previous_reading), _dec(0)),
                reset_check("WaterReadings", n, idx, values, rec, "current_reading", _dec(rec.current_reading), _dec(0)),
                reset_check("WaterReadings", n, idx, values, rec, "rate", _dec(rec.rate), _dec(db.water_rate)),
                reset_check("WaterReadings", n, idx, values, rec, "reading_date", rec.reading_date, date.today()),
                reset_check("WaterReadings", n, idx, values, rec, "paid", bool(rec.paid), False),
                differs(prev, _dec(rec.previous_reading)), differs(cur, _dec(rec.current_reading)),
                differs(rate, _dec(rec.rate)), differs(rdate, rec.reading_date), differs(pdate, rec.paid_date),
                differs(paid, bool(rec.paid))])
        if month and closed(month) and changed:
            report.add("WaterReadings", n, "reading_month", "PERIOD_CLOSED", "blocked")
        if ref and month:
            seen.setdefault((get(idx, values, "unit_id"), month), []).append(n)
            proposal[n] = "insert"
            if ref[0] == "db" and (ref[1].id, month) in db.readings:
                found = db.readings[(ref[1].id, month)]
                if db.crosswalk.get(("water_reading", ext)) != found.id:
                    report.add("WaterReadings", n, "reading_month", "UNIT_MONTH_EXISTS", "conflict")
                else:
                    proposal[n] = "update" if changed else "unchanged"
    for ns in seen.values():
        if len(ns) > 1:
            for n in ns:
                report.add("WaterReadings", n, "reading_month", "UNIT_MONTH_DUPLICATE", "blocked")
    for n, _ in rows:
        if report.row_status("WaterReadings", n) == "valid" and source_ok and n in proposal:
            report.propose("water_reading", proposal[n])

    # ---------- Billing
    idx, rows = table("Billing")
    pidx, prow = table("Payments")
    bills = {}           # workbook bill id -> info
    seen = {}
    billing_total, payment_total = Decimal("0"), Decimal("0")
    first_month_of_unit = {}
    dup = check_ext_column(idx, rows)
    for n, values in rows:
        report.see("Billing", n)
        ext = ext_of(idx, values)
        if ext is None:
            report.add("Billing", n, "id", "EXTERNAL_ID_MISSING", "blocked")
        elif n in dup:
            report.add("Billing", n, "id", "EXTERNAL_ID_DUPLICATE", "blocked")
        rec = existing("Billing", n, "billing", ext)
        ref = resolve_unit("Billing", n, idx, values)
        month = a_month("Billing", n, "billing_month", idx, values)
        points_elsewhere("Billing", n, "billing_month", rec, ref, extra=(rec is None or rec.billing_month == month))
        charges = {c: number("Billing", n, c, idx, values, negative=(c == "adjustment")) for c in BILL_CHARGES}
        prev_bal = number("Billing", n, "previous_balance", idx, values)
        paid = number("Billing", n, "amount_paid", idx, values)
        due = a_date("Billing", n, "due_date", idx, values)
        paid_date = a_date("Billing", n, "paid_date", idx, values)
        status = choice("Billing", n, "status", idx, values, BILL_STATUSES, "STATUS_INVALID", current=getattr(rec, "status", None))
        changed = rec is None or any(
            [differs(charges[c], _dec(getattr(rec, c))) for c in BILL_CHARGES]
            + [differs(prev_bal, _dec(rec.previous_balance)), differs(paid, _dec(rec.amount_paid)),
               differs(status, rec.status), differs(due, rec.due_date), differs(paid_date, rec.paid_date)])
        charge_sum = sum((v for v in charges.values() if v is not None), Decimal("0"))
        billing_total += charge_sum
        unit_key = get(idx, values, "unit_id")
        if month:
            seen.setdefault((unit_key, month), []).append(n)
            if closed(month) and changed:
                report.add("Billing", n, "billing_month", "PERIOD_CLOSED", "blocked")
            if cutoff_month and month > cutoff_month:
                report.add("Billing", n, "billing_month", "BILL_AFTER_CUTOFF", "blocked")
            if (first_m and month < first_m) or (last_m and month > last_m):
                report.add("Billing", n, "billing_month", "BILL_OUTSIDE_SCOPE", "blocked")
            if mode == "opening_balances" and cutoff_month and month < cutoff_month:
                report.add("Billing", n, "billing_month", "HISTORY_IN_OPENING_BALANCE_MODE", "conflict")
            fm = first_month_of_unit.get(unit_key)
            first_month_of_unit[unit_key] = month if fm is None or month < fm else fm
        if ref and month and ref[0] == "db" and (ref[1].id, month) in db.bills:
            found = db.bills[(ref[1].id, month)]
            if db.crosswalk.get(("billing", ext)) != found.id:
                report.add("Billing", n, "billing_month", "UNIT_MONTH_EXISTS", "conflict")
            if changed and getattr(found, "soa_manual_override", False):
                report.add("Billing", n, "billing_month", "BILL_CORRECTED_BY_HAND", "blocked")
            if changed and getattr(found, "issued_at", None) is not None:
                report.add("Billing", n, "billing_month", "BILL_ISSUED", "blocked")
        if charges.get("penalty") and changed:
            report.add("Billing", n, "penalty", "PENALTY_RECALCULATED", "unresolved")
        if ext:
            bills[ext] = {"row": n, "paid": paid, "charges": charge_sum, "prev": prev_bal, "status": status,
                          "unit": unit_key, "month": month, "valid_ref": ref is not None, "changed": changed}
    for ns in seen.values():
        if len(ns) > 1:
            for n in ns:
                report.add("Billing", n, "billing_month", "UNIT_MONTH_DUPLICATE", "blocked")
    # carried balances: CITYLAND9 recalculates arrears from earlier bills (_prepare_bill_calculation_cache)
    for info in bills.values():
        if info["changed"] and info["prev"] and info["prev"] > 0 and info["month"]:
            if info["month"] == first_month_of_unit.get(info["unit"]):
                report.add("Billing", info["row"], "previous_balance", "OPENING_ARREARS_NOT_CARRIED", "unresolved")
            else:
                report.add("Billing", info["row"], "previous_balance", "PREVIOUS_BALANCE_WITH_HISTORY", "unresolved")

    # ---------- Payments
    paid_by_bill, payment_parent, payment_changed = {}, {}, set()
    dup = check_ext_column(pidx, prow)
    for n, values in prow:
        report.see("Payments", n)
        ext = ext_of(pidx, values)
        if ext is None:
            report.add("Payments", n, "id", "EXTERNAL_ID_MISSING", "blocked")
        elif n in dup:
            report.add("Payments", n, "id", "EXTERNAL_ID_DUPLICATE", "blocked")
        rec = existing("Payments", n, "payment", ext)
        ref = external_id_text(get(pidx, values, "billing_id"))
        amount = number("Payments", n, "amount", pidx, values, positive=True, required=True)
        d = a_date("Payments", n, "payment_date", pidx, values, required=(rec is None))
        method = choice("Payments", n, "payment_method", pidx, values, PAYMENT_METHODS, "STATUS_INVALID",
                        current=getattr(rec, "payment_method", None))
        if ref is None or (ref not in bills and ("billing", ref) not in db.crosswalk):
            report.add("Payments", n, "billing_id", "UNRESOLVED_BILL_REFERENCE", "blocked")
        else:
            if rec is not None and db.crosswalk.get(("billing", ref)) != rec.billing_id:
                report.add("Payments", n, "billing_id", "ID_POINTS_TO_OTHER_RECORD", "conflict")
            if ref in bills and amount is not None:
                paid_by_bill[ref] = paid_by_bill.get(ref, Decimal("0")) + amount
        if amount is not None:
            payment_total += amount
        changed = rec is None or any([differs(amount, _dec(rec.amount)), differs(d, rec.payment_date),
                                      differs(method, rec.payment_method),
                                      differs(text_field("Payments", n, "reference", pidx, values, 100), rec.reference),
                                      differs(text_field("Payments", n, "remarks", pidx, values, 300), rec.remarks)])
        if changed:
            payment_changed.add(n)
        if d and closed(d.strftime("%Y-%m")) and changed:
            report.add("Payments", n, "payment_date", "PERIOD_CLOSED", "blocked")
        payment_parent[n] = ref
    for ext, info in bills.items():
        touched = info["changed"] or any(payment_parent.get(n) == ext for n in payment_changed)
        if not touched:
            continue      # an unchanged bill and its unchanged payments: the import would not alter them
        paid = info["paid"] or Decimal("0")
        detail = paid_by_bill.get(ext, Decimal("0"))
        if paid > 0 and ext not in paid_by_bill:
            report.add("Billing", info["row"], "amount_paid", "PAYMENT_DETAIL_MISSING", "unresolved")
        elif paid != detail:
            report.add("Billing", info["row"], "amount_paid", "PAYMENT_TOTAL_MISMATCH", "conflict")
        owed = info["charges"] + (info["prev"] or Decimal("0"))
        if paid > owed and owed >= 0:
            report.add("Billing", info["row"], "amount_paid", "OVERPAYMENT_WITHOUT_ADVANCE_LEDGER", "unresolved")
        if info["status"] == "Paid" and paid < owed:
            report.add("Billing", info["row"], "status", "STATUS_INCONSISTENT", "conflict")
        if info["status"] in ("Unpaid", "Overdue") and owed > 0 and paid >= owed:
            report.add("Billing", info["row"], "status", "STATUS_INCONSISTENT", "conflict")
    for n, _ in rows:
        if report.row_status("Billing", n) == "valid" and source_ok:
            info = next((i for i in bills.values() if i["row"] == n), None)
            report.propose("billing", "unchanged" if info and not info["changed"] else ("update" if rt else "insert"))
    for n, ref in payment_parent.items():   # after the bill checks: a payment can't go ahead of its bill
        if ref in bills and report.row_status("Billing", bills[ref]["row"]) != "valid":
            report.add("Payments", n, "billing_id", "PARENT_RECORD_NOT_VALID", "unresolved")
        if report.row_status("Payments", n) == "valid" and source_ok:
            report.propose("payment", ("update" if n in payment_changed else "unchanged") if rt else "insert")

    # ---------- Employees
    idx, rows = table("Employees")
    nos = {}
    for n, values in rows:
        no = str(get(idx, values, "employee_no") or "").strip()
        if no:
            nos.setdefault(norm_key(no), []).append(n)
    for n, values in rows:
        report.see("Employees", n)
        ext = ext_of(idx, values)
        existing("Employees", n, "employee", ext)
        no = text_field("Employees", n, "employee_no", idx, values, 50)
        if not no:
            report.add("Employees", n, "employee_no", "EMPLOYEE_NO_MISSING", "blocked")
        elif len(nos.get(norm_key(no), [])) > 1:
            report.add("Employees", n, "employee_no", "EMPLOYEE_NO_DUPLICATE", "blocked")
        text_field("Employees", n, "full_name", idx, values, 200, required=True)
        a_date("Employees", n, "date_hired", idx, values)
        number("Employees", n, "monthly_salary", idx, values)
        by_no = db.employees_by_no.get(norm_key(no)) if no else None
        mapped = db.crosswalk.get(("employee", ext)) if ext else None
        if (mapped is not None or rt) and by_no is not None and by_no.id != mapped:
            report.add("Employees", n, "employee_no", "EMPLOYEE_NO_CONFLICT", "conflict")
        if report.row_status("Employees", n) == "valid" and source_ok:
            report.propose("employee", "update" if (mapped is not None or by_no is not None) else "insert")

    # ---------- Expenses
    idx, rows = table("Expenses")
    for n, values in rows:
        report.see("Expenses", n)
        ext = ext_of(idx, values)
        rec = existing("Expenses", n, "expense", ext)
        d = a_date("Expenses", n, "expense_date", idx, values, required=(rec is None))
        category = text_field("Expenses", n, "category", idx, values, 100, required=(rec is None))
        description = text_field("Expenses", n, "description", idx, values, 300)
        amount = number("Expenses", n, "amount", idx, values, positive=True, required=(rec is None))
        changed = rec is None
        if rec is not None:
            changed = any([
                reset_check("Expenses", n, idx, values, rec, "expense_date", rec.expense_date, date.today()),
                reset_check("Expenses", n, idx, values, rec, "category", rec.category, "OTHERS"),
                reset_check("Expenses", n, idx, values, rec, "description", rec.description or "", ""),
                reset_check("Expenses", n, idx, values, rec, "amount", _dec(rec.amount), _dec(0)),
                differs(d, rec.expense_date), differs(category, rec.category),
                differs(description, rec.description), differs(amount, _dec(rec.amount))])
        if d and closed(d.strftime("%Y-%m")) and changed:
            report.add("Expenses", n, "expense_date", "PERIOD_CLOSED", "blocked")
        if not rt and (ext is None or ("expense", ext) not in db.crosswalk):
            report.add("Expenses", n, "id", "NO_NATURAL_KEY", "warning")
        if report.row_status("Expenses", n) == "valid" and source_ok:
            report.propose("expense", ("update" if changed else "unchanged") if rt else "insert")

    # ---------- readiness
    for i in report.issues:
        if i["severity"] in ("blocked", "conflict", "unresolved"):
            area = "financial" if (i["sheet"] in ("Billing", "Payments") and not rt) else "records"
            report.block_area(area, i["code"])
    report.readiness.setdefault("records", [])
    if rt:
        return
    report.block_area("financial", "ADVANCE_LEDGER_MISSING")      # this format has no advances/credits
    report.block_area("financial", "RECEIPTS_NOT_IN_SOURCE")      # nor official receipts
    if mode == "opening_balances":
        report.block_area("financial", "OPENING_BALANCE_UNSUPPORTED")
    totals = fin.get("control_totals") or {}
    if all(totals.get(k) not in (None, "") for k in ("billing_rows", "billing_charges_total", "payment_rows", "payment_total")):
        try:
            matches = (int(totals["billing_rows"]) == len(table("Billing")[1])
                       and Decimal(str(totals["billing_charges_total"])) == billing_total
                       and int(totals["payment_rows"]) == len(prow)
                       and Decimal(str(totals["payment_total"])) == payment_total)
        except Exception:
            matches = False
        if not matches:
            report.block_area("financial", "CONTROL_TOTALS_MISMATCH")
        report.notes.append(f"Control totals {'match' if matches else 'do NOT match'} the workbook.")


# Unresolved findings that still stop a round-trip import (they would create a duplicate or orphan).
GATE_UNRESOLVED = ("POSSIBLE_DUPLICATE_PERSON", "PARENT_RECORD_NOT_VALID")


def import_gate(report):
    """May the classic import (run_excel_import) run on this workbook? Only CITYLAND9's own export of
    THIS database, every row updating an existing record safely (cityland_roundtrip check)."""
    if report.get("format") != "cityland_export" or report.get("mode") != "cityland_roundtrip":
        return {"allowed": False, "blockers": [{"code": "LAYOUT_NOT_IMPORTABLE", "message": MESSAGES["LAYOUT_NOT_IMPORTABLE"], "count": 1}]}
    blockers = [{"code": code, "message": c["message"], "count": c["count"]} for code, c in report["issueCounts"].items()
                if c["severity"] in ("blocked", "conflict") or code in GATE_UNRESOLVED]
    if not report["sheets"]:
        blockers.append({"code": "FORMAT_UNKNOWN", "message": MESSAGES["FORMAT_UNKNOWN"], "count": 1})
    return {"allowed": not blockers, "blockers": blockers}


# ------------------------------------------------------------------ entry point
def detect_format(sheets, config):
    prop_sheet = (((config or {}).get("properties") or {}).get("sheet")) or "Properties"
    if prop_sheet in sheets:
        return "properties"
    if any(name in sheets for name in EXPORT_SHEETS):
        return "cityland_export"
    return "unknown"


def check_workbook(stream, filename, *, legacy, config, roundtrip="auto"):
    """Validate a workbook against CITYLAND9 without writing anything. Returns the report dict.
    roundtrip: "auto" = CITYLAND9's own export layout is checked as an edit of THIS database
    (its ids are this database's ids); True/False force it. Other layouts are migrations (check only).
    The report carries "importGate": may the classic import run on this file?
    Raises CheckError when the file can't be read at all."""
    stream.seek(0)
    data = stream.read()
    file_info = {"name": os.path.basename(filename or "workbook.xlsx")[:255], "sha256": hashlib.sha256(data).hexdigest(),
                 "bytes": len(data)}
    import io
    sheets, limit_problems = read_workbook(io.BytesIO(data), max_unpacked_mb=legacy.IMPORT_MAX_UNCOMPRESSED_MB,
                                           max_rows=legacy.IMPORT_MAX_ROWS, max_columns=legacy.IMPORT_MAX_COLUMNS)
    report = Report()
    for sheet, code in limit_problems:
        report.add(sheet, None, None, code, "blocked")
        report.block_area("file", code)
    fmt = detect_format(sheets, config)
    rt = fmt == "cityland_export" and (roundtrip is True or roundtrip == "auto")
    config_problems = []
    if not rt and not source_code_of(config):
        config_problems.append("SOURCE_SYSTEM_NOT_CONFIGURED")
    with no_writes(legacy.db):
        db = DbView(legacy, source_code_of(config), roundtrip=rt)
        if not db.crosswalk_ok:
            config_problems.append("CROSSWALK_UNAVAILABLE")
        if fmt == "properties":
            check_properties(sheets, config, db, report)
        elif fmt == "cityland_export":
            check_export(sheets, config, db, report)
        else:
            report.add(None, None, None, "FORMAT_UNKNOWN", "blocked")
            report.block_area("file", "FORMAT_UNKNOWN")
    if not rt:
        report.block_area("import", "IMPORTER_NOT_ENABLED")
    result = report.as_dict(fmt=fmt, file_info=file_info, config_problems=config_problems)
    result["mode"] = "cityland_roundtrip" if rt else "migration"
    result["importGate"] = import_gate(result)
    if rt:
        result["readiness"]["import"] = {"ready": result["importGate"]["allowed"],
                                         "blockers": [{"code": b["code"], "message": b["message"]} for b in result["importGate"]["blockers"]]}
        result["readyForImport"] = result["importGate"]["allowed"]
    return result
